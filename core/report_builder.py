"""Report(회차 보고서)를 실제 산출물과 동일한 7섹션 PDF로 렌더링하는 모듈 (reportlab).

레이아웃은 실제로 다운로드받은 예시 PDF(건설재해예방전문지도기관 기술지도 결과보고서)를
그대로 기준으로 삼았다. 한글 폰트는 Windows 기본 맑은 고딕(malgun.ttf)을 사용한다 —
이 앱은 Windows 데스크톱 배포를 전제로 하므로 별도 폰트 파일을 동봉하지 않았다.
"""

from __future__ import annotations

from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from core.constants import FIXED_HAZARD_FACTORS, MEASUREMENT_INSTRUMENTS
from core.db import SessionLocal
from core.models_db import Report

FONT_REGULAR = "MalgunGothic"
FONT_BOLD = "MalgunGothic-Bold"
_FONTS_REGISTERED = False

_WINDOWS_FONT_DIR = Path(r"C:\Windows\Fonts")

_HEADER_BG = colors.HexColor("#f3f4f6")
_ACCENT = colors.HexColor("#4f46e5")


def _ensure_fonts() -> None:
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return
    regular = _WINDOWS_FONT_DIR / "malgun.ttf"
    bold = _WINDOWS_FONT_DIR / "malgunbd.ttf"
    if not regular.exists():
        raise RuntimeError(
            "맑은 고딕 폰트를 찾을 수 없습니다 (C:\\Windows\\Fonts\\malgun.ttf). "
            "Windows 환경에서 실행 중인지 확인하세요."
        )
    pdfmetrics.registerFont(TTFont(FONT_REGULAR, str(regular)))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, str(bold) if bold.exists() else str(regular)))
    _FONTS_REGISTERED = True


def _p(text: str, bold: bool = False, size: int = 9, align: str = "LEFT", color=colors.black) -> Paragraph:
    style = ParagraphStyle(
        name="cell",
        fontName=FONT_BOLD if bold else FONT_REGULAR,
        fontSize=size,
        leading=size + 4,
        alignment={"LEFT": 0, "CENTER": 1, "RIGHT": 2}[align],
        textColor=color,
    )
    return Paragraph((text or "-").replace("\n", "<br/>"), style)


def _scaled_image(path: str, max_width: float, max_height: float):
    if not path or not Path(path).exists():
        return _p("(사진 없음)", align="CENTER")
    try:
        img = Image(path)
        ratio = min(max_width / img.imageWidth, max_height / img.imageHeight, 1.0)
        img.drawWidth = img.imageWidth * ratio
        img.drawHeight = img.imageHeight * ratio
        return img
    except Exception:
        return _p("(이미지를 불러올 수 없음)", align="CENTER")


_GRID = TableStyle(
    [
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#9ca3af")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME", (0, 0), (-1, -1), FONT_REGULAR),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]
)


def _section_title(number: int, text: str) -> Paragraph:
    style = ParagraphStyle(
        name="section",
        fontName=FONT_BOLD,
        fontSize=12,
        spaceAfter=4,
        spaceBefore=10,
    )
    return Paragraph(f"{number}. {text}", style)


def _fmt_date(value) -> str:
    return value.strftime("%Y.%m.%d") if value else "-"


def _fmt_amount(value: int | None) -> str:
    return f"{value:,}" if value is not None else "-"


def _build_site_section(site) -> list:
    label = lambda t: _p(t, bold=True, align="CENTER")
    table = Table(
        [
            [label("현장"), label("현장명"), _p(site.name), label("사업장관리번호\n(사업개시번호)"), _p(f"{site.site_mgmt_no}\n{site.biz_start_no}")],
            ["", label("공사기간"), _p(f"{_fmt_date(site.period_start)} ~ {_fmt_date(site.period_end)}"), label("공사금액"), _p(_fmt_amount(site.amount))],
            ["", label("책임자"), _p(site.manager_name), label("연락처\n(이메일)"), _p(f"{site.manager_phone}\n{site.manager_email}")],
            ["", label("주소"), _p(site.address, align="LEFT"), "", ""],
            [label("본사"), label("회사명"), _p(site.hq_company), label("법인등록번호\n(사업자등록번호)"), _p(f"{site.corp_reg_no}\n{site.biz_reg_no}")],
            ["", label("면허번호"), _p(site.license_no), label("연락처"), _p(site.hq_phone)],
            ["", label("주소"), _p(site.hq_address, align="LEFT"), "", ""],
        ],
        colWidths=[16 * mm, 26 * mm, 55 * mm, 30 * mm, 43 * mm],
    )
    style = TableStyle(_GRID.getCommands())
    style.add("SPAN", (0, 0), (0, 3))
    style.add("SPAN", (0, 4), (0, 6))
    style.add("SPAN", (2, 3), (4, 3))
    style.add("SPAN", (2, 6), (4, 6))
    style.add("BACKGROUND", (0, 0), (0, -1), _HEADER_BG)
    style.add("BACKGROUND", (1, 0), (1, -1), _HEADER_BG)
    style.add("BACKGROUND", (3, 0), (3, -1), _HEADER_BG)
    table.setStyle(style)
    return [_section_title(1, "기술지도 대상사업장"), table]


def _build_overview_meta_section(report: Report, site) -> list:
    label = lambda t: _p(t, bold=True, align="CENTER")
    guidance_type = "■ 건설 / □ 전기·정보통신"
    implemented = "■ 이행 / □ 불이행" if report.prev_guidance_implemented else "□ 이행 / ■ 불이행"
    if report.prev_guidance_implemented is None:
        implemented = "□ 이행 / □ 불이행"

    methods = ["직접전달", "등기우편", "전자우편", "모바일", "기타"]
    method_text = " / ".join(
        f"{'■' if m == report.notification_method else '□'} {m}" for m in methods
    )

    staff_name = report.assigned_staff.name if report.assigned_staff else "-"
    staff_phone = report.assigned_staff.phone if report.assigned_staff else "-"

    table = Table(
        [
            [label("지도기관명"), _p("(주)한국미래안전"), label("기술지도실시일"), _p(_fmt_date(report.guidance_date))],
            [label("구분"), _p(guidance_type), label("공정률"), _p(f"({report.progress_rate or 0})%")],
            [label("횟수"), _p(f"({report.visit_no})회차 / 총({site.total_guidance_count or '-'})회"), label("담당 요원"), _p(staff_name)],
            [label("이전 기술지도\n이행여부"), _p(implemented), label("연락처"), _p(staff_phone)],
            [label("현장책임자 등\n통보방법"), _p(method_text, align="LEFT"), "", ""],
            [label("기타 특이사항"), _p(report.special_note, align="LEFT"), "", ""],
        ],
        colWidths=[26 * mm, 71 * mm, 26 * mm, 47 * mm],
    )
    style = TableStyle(_GRID.getCommands())
    style.add("SPAN", (1, 4), (3, 4))
    style.add("SPAN", (1, 5), (3, 5))
    style.add("BACKGROUND", (0, 0), (0, -1), _HEADER_BG)
    style.add("BACKGROUND", (2, 0), (2, 3), _HEADER_BG)
    table.setStyle(style)
    return [_section_title(2, "기술지도 개요"), table]


def _build_photo_and_hazard_section(report: Report) -> list:
    label = lambda t: _p(t, bold=True, align="CENTER")
    photos = {p.slot: p.photo_path for p in report.overview_photos}
    photo_row = [
        _scaled_image(photos.get(1, ""), 60 * mm, 45 * mm),
        _scaled_image(photos.get(2, ""), 60 * mm, 45 * mm),
    ]
    overview_table = Table(
        [[label("위치"), label("전경 1"), label("전경 2")], [label("사진"), *photo_row]],
        colWidths=[16 * mm, 70 * mm, 70 * mm],
        rowHeights=[8 * mm, 50 * mm],
    )
    overview_table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))

    checked = set(report.hazard_factor_checks or [])
    hazard_rows = [[label("구분"), label("12대 기인물"), label("안전조치")]]
    for number, name, action in FIXED_HAZARD_FACTORS:
        mark = "■" if number in checked else "□"
        hazard_rows.append([_p(mark, align="CENTER"), _p(f"{number}. {name}"), _p(action)])
    hazard_table = Table(hazard_rows, colWidths=[16 * mm, 40 * mm, 100 * mm])
    hazard_table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))

    return [
        _section_title(3, "현장 전경사진"),
        overview_table,
        Spacer(1, 8 * mm),
        _p("건설현장 12대 사망사고 기인물 안전조치", bold=True, size=11),
        Spacer(1, 2 * mm),
        hazard_table,
    ]


def _build_previous_findings_section(report: Report) -> list:
    label = lambda t: _p(t, bold=True, align="CENTER")
    rows = [[label("지도일\n(확인일)"), label("유해·위험요인/\n지적사진"), label("지적사항"), label("조치내용"), label("이행결과")]]
    previous = {pf.slot: pf for pf in report.previous_findings}
    for slot in range(1, 5):
        pf = previous.get(slot)
        if pf:
            photo_cell = _scaled_image(pf.photo_path, 28 * mm, 20 * mm) if pf.photo_path else _p(pf.title, size=8)
            rows.append([_p(_fmt_date(report.guidance_date), size=8), photo_cell, _p(pf.content, size=8), _p(pf.action_result, size=8), _p("이행" if pf.confirmed else "-", size=8)])
        else:
            rows.append(["", "", "", "", ""])
    table = Table(rows, colWidths=[20 * mm, 32 * mm, 60 * mm, 30 * mm, 20 * mm], rowHeights=[10 * mm] + [22 * mm] * 4)
    table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))
    return [_section_title(4, "이전 기술지도 사항 이행여부"), table]


def _build_current_risk_section(report: Report) -> list:
    label = lambda t: _p(t, bold=True, align="CENTER")
    rows = [[label("유해·위험요인\n지적사진"), label("지적사항(개선대책)"), label("비고")]]
    findings = {f.slot: f for f in report.findings}
    for slot in range(1, 5):
        f = findings.get(slot)
        if f:
            photo_cell = _scaled_image(f.photo_path, 28 * mm, 20 * mm)
            content = f"{f.title}\n{f.content}"
            if f.law_citation:
                content += f"\n* {f.law_citation}"
            risk_text = ""
            if f.risk_score is not None:
                risk_text = f"가능성 {f.likelihood} / 중대성 {f.severity}\n위험성 {f.risk_score}"
            rows.append([photo_cell, _p(content, size=8), _p(risk_text, size=8, align="CENTER")])
        else:
            rows.append(["", "", ""])
    table = Table(rows, colWidths=[32 * mm, 100 * mm, 30 * mm], rowHeights=[10 * mm] + [26 * mm] * 4)
    table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))
    return [_section_title(5, "현재 공정 내 현존하는 위험성 제거"), table]


def _build_process_section(report: Report) -> list:
    label = lambda t: _p(t, bold=True, align="CENTER")
    matrix_rows = [
        [label("빈도/강도"), label("1"), label("2"), label("3")],
        [label("1"), _p("1", align="CENTER"), _p("2", align="CENTER"), _p("3", align="CENTER")],
        [label("2"), _p("2", align="CENTER"), _p("4", align="CENTER"), _p("6", align="CENTER")],
        [label("3"), _p("3", align="CENTER"), _p("6", align="CENTER"), _p("9", align="CENTER")],
    ]
    matrix_table = Table(matrix_rows, colWidths=[18 * mm, 14 * mm, 14 * mm, 14 * mm])
    matrix_table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG), ("BACKGROUND", (0, 0), (0, -1), _HEADER_BG)]))

    level_rows = [
        [label("위험성 수준"), label("관리기준")],
        [_p("1~2 (하)", align="CENTER"), _p("현재상태유지", align="CENTER")],
        [_p("3~4 (중)", align="CENTER"), _p("개선", align="CENTER")],
        [_p("6~9 (상)", align="CENTER"), _p("즉시개선", align="CENTER")],
    ]
    level_table = Table(level_rows, colWidths=[35 * mm, 35 * mm])
    level_table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))

    header_row = Table([[matrix_table, level_table]], colWidths=[65 * mm, 95 * mm])
    header_row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))

    rows = [[label("진행공정"), label("유해·위험요인"), label("예방대책"), label("위험성수준")]]
    for entry in report.process_entries:
        rows.append([_p(entry.process_name, size=8), _p(entry.hazard_text, size=8), _p(entry.prevention_text, size=8), _p(entry.risk_level, size=8, align="CENTER")])
    while len(rows) < 5:
        rows.append(["", "", "", ""])
    process_table = Table(rows, colWidths=[30 * mm, 55 * mm, 55 * mm, 22 * mm], rowHeights=[10 * mm] + [20 * mm] * 4)
    process_table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))

    return [
        _section_title(6, "향후 진행공정에 대한 유해·위험 요인 파악 및 대책"),
        header_row,
        Spacer(1, 4 * mm),
        process_table,
    ]


def _build_support_section(report: Report) -> list:
    label = lambda t: _p(t, bold=True, align="CENTER")
    education = report.safety_education
    materials = report.provided_materials

    education_text = f"○ 참석인원: {education.attendee_count or '-'}명" if education else "○ 참석인원: -"
    material_names = "\n".join(f"- {m.title}" for m in materials) or "-"

    measurement_by_type = {m.instrument_type: m for m in report.measurements}
    measurement_rows = [[label("장비명"), label("측정값"), label("내용")]]
    for name, unit in MEASUREMENT_INSTRUMENTS:
        m = measurement_by_type.get(name)
        value = f"{m.value} {unit}" if m and m.value else "-"
        measurement_rows.append([_p(name, size=8), _p(value, size=8, align="CENTER"), ""])
    measurement_table = Table(measurement_rows, colWidths=[38 * mm, 30 * mm, 42 * mm])
    measurement_table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))

    rows = [
        [label("지원사항"), label("구체적 사항"), label("비고")],
        [_p("○ 교육"), _p(education_text), ""],
        [_p("○ 자료배포"), _p(material_names), ""],
        [_p("○ 장비 측정"), measurement_table, ""],
    ]
    table = Table(rows, colWidths=[20 * mm, 116 * mm, 26 * mm])
    table.setStyle(TableStyle(_GRID.getCommands() + [("BACKGROUND", (0, 0), (-1, 0), _HEADER_BG)]))
    return [_section_title(7, "사업장 지원 사항 등 기타 사항"), table]


def _build_material_appendix(report: Report) -> list:
    flowables: list = []
    for material in report.provided_materials:
        source_path = material.custom_photo_path
        if not source_path and material.material:
            source_path = material.material.file_path
        if source_path and Path(source_path).suffix.lower() in (".jpg", ".jpeg", ".png", ".webp"):
            flowables.append(PageBreak())
            flowables.append(_scaled_image(source_path, 170 * mm, 250 * mm))
    return flowables


def build_report(report_id: int, output_path: str | Path) -> Path:
    """Report 하나를 실제 산출물과 동일한 구조의 PDF로 생성한다."""
    _ensure_fonts()

    with SessionLocal() as session:
        report = session.get(Report, report_id)
        if report is None:
            raise ValueError(f"Report {report_id}를 찾을 수 없습니다.")
        site = report.site

        story: list = []
        title_style = ParagraphStyle(name="title", fontName=FONT_BOLD, fontSize=16, alignment=1, spaceAfter=10)
        story.append(Paragraph("건설재해예방전문지도기관 기술지도 결과보고서", title_style))

        story += _build_site_section(site)
        story.append(Spacer(1, 4 * mm))
        story += _build_overview_meta_section(report, site)

        story.append(PageBreak())
        story += _build_photo_and_hazard_section(report)

        story.append(PageBreak())
        story += _build_previous_findings_section(report)

        story.append(PageBreak())
        story += _build_current_risk_section(report)

        story.append(PageBreak())
        story += _build_process_section(report)

        story.append(PageBreak())
        story += _build_support_section(report)

        story += _build_material_appendix(report)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=15 * mm,
            rightMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
        )
        doc.build(story)

        report.pdf_path = str(output_path)
        report.status = "final"
        session.commit()

    return output_path


# ---------------------------------------------------------------------------
# DOCX (python-docx) — PDF와 같은 7섹션을 워드 문서로도 만든다.
# HWPX는 core/hwpx_exporter.py가 이 DOCX를 한글에서 열어 다른 이름으로 저장하는
# 방식으로 만들기 때문에, DOCX 레이아웃이 정확할수록 HWPX 품질도 좋아진다.
# ---------------------------------------------------------------------------


def _docx_set_cell(cell, text: str, bold: bool = False, align: str = "left") -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    for i, line in enumerate((text or "-").split("\n")):
        if i > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold


def _docx_add_table(document, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    return table


def build_report_docx(report_id: int, output_path: str | Path) -> Path:
    """Report 하나를 PDF와 동일한 구조의 DOCX로 생성한다."""
    from docx import Document
    from docx.shared import Pt

    with SessionLocal() as session:
        report = session.get(Report, report_id)
        if report is None:
            raise ValueError(f"Report {report_id}를 찾을 수 없습니다.")
        site = report.site

        document = Document()
        style = document.styles["Normal"]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(9)

        document.add_heading("건설재해예방전문지도기관 기술지도 결과보고서", level=1)

        # 1. 기술지도 대상사업장
        document.add_heading("1. 기술지도 대상사업장", level=2)
        t = _docx_add_table(document, 7, 5)
        rows_data = [
            ("현장", "현장명", site.name, "사업장관리번호\n(사업개시번호)", f"{site.site_mgmt_no}\n{site.biz_start_no}"),
            ("", "공사기간", f"{_fmt_date(site.period_start)} ~ {_fmt_date(site.period_end)}", "공사금액", _fmt_amount(site.amount)),
            ("", "책임자", site.manager_name, "연락처(이메일)", f"{site.manager_phone}\n{site.manager_email}"),
            ("", "주소", site.address, "", ""),
            ("본사", "회사명", site.hq_company, "법인등록번호\n(사업자등록번호)", f"{site.corp_reg_no}\n{site.biz_reg_no}"),
            ("", "면허번호", site.license_no, "연락처", site.hq_phone),
            ("", "주소", site.hq_address, "", ""),
        ]
        for r, row in enumerate(rows_data):
            for c, value in enumerate(row):
                _docx_set_cell(t.cell(r, c), value, bold=(c in (0, 1, 3)))

        # 2. 기술지도 개요
        document.add_heading("2. 기술지도 개요", level=2)
        methods = ["직접전달", "등기우편", "전자우편", "모바일", "기타"]
        method_text = " / ".join(f"{'■' if m == report.notification_method else '□'} {m}" for m in methods)
        implemented = "□ 이행 / □ 불이행"
        if report.prev_guidance_implemented is True:
            implemented = "■ 이행 / □ 불이행"
        elif report.prev_guidance_implemented is False:
            implemented = "□ 이행 / ■ 불이행"
        staff_name = report.assigned_staff.name if report.assigned_staff else "-"
        staff_phone = report.assigned_staff.phone if report.assigned_staff else "-"

        t2 = _docx_add_table(document, 6, 4)
        rows2 = [
            ("지도기관명", "(주)한국미래안전", "기술지도실시일", _fmt_date(report.guidance_date)),
            ("구분", "■ 건설 / □ 전기·정보통신", "공정률", f"({report.progress_rate or 0})%"),
            ("횟수", f"({report.visit_no})회차 / 총({site.total_guidance_count or '-'})회", "담당 요원", staff_name),
            ("이전 기술지도\n이행여부", implemented, "연락처", staff_phone),
            ("현장책임자 등\n통보방법", method_text, "", ""),
            ("기타 특이사항", report.special_note, "", ""),
        ]
        for r, row in enumerate(rows2):
            for c, value in enumerate(row):
                _docx_set_cell(t2.cell(r, c), value, bold=(c in (0, 2)))

        document.add_page_break()

        # 3. 현장 전경사진 + 12대 기인물
        document.add_heading("3. 현장 전경사진", level=2)
        overview_table = _docx_add_table(document, 2, 3)
        _docx_set_cell(overview_table.cell(0, 0), "위치", bold=True)
        _docx_set_cell(overview_table.cell(0, 1), "전경 1", bold=True)
        _docx_set_cell(overview_table.cell(0, 2), "전경 2", bold=True)
        _docx_set_cell(overview_table.cell(1, 0), "사진", bold=True)
        photos = {p.slot: p.photo_path for p in report.overview_photos}
        for slot, col in ((1, 1), (2, 2)):
            path = photos.get(slot, "")
            cell = overview_table.cell(1, col)
            cell.text = ""
            if path and Path(path).exists():
                try:
                    cell.paragraphs[0].add_run().add_picture(path, width=Pt(120))
                except Exception:
                    _docx_set_cell(cell, "(이미지를 불러올 수 없음)")
            else:
                _docx_set_cell(cell, "(사진 없음)")

        document.add_paragraph()
        document.add_paragraph("건설현장 12대 사망사고 기인물 안전조치").runs[0].bold = True
        checked = set(report.hazard_factor_checks or [])
        hazard_table = _docx_add_table(document, 1 + len(FIXED_HAZARD_FACTORS), 3)
        for c, text in enumerate(("구분", "12대 기인물", "안전조치")):
            _docx_set_cell(hazard_table.cell(0, c), text, bold=True)
        for i, (number, name, action) in enumerate(FIXED_HAZARD_FACTORS, start=1):
            mark = "■" if number in checked else "□"
            _docx_set_cell(hazard_table.cell(i, 0), mark, align="center")
            _docx_set_cell(hazard_table.cell(i, 1), f"{number}. {name}")
            _docx_set_cell(hazard_table.cell(i, 2), action)

        document.add_page_break()

        # 4. 이전 기술지도 사항 이행여부
        document.add_heading("4. 이전 기술지도 사항 이행여부", level=2)
        prev_table = _docx_add_table(document, 5, 5)
        for c, text in enumerate(("지도일(확인일)", "유해·위험요인/지적사진", "지적사항", "조치내용", "이행결과")):
            _docx_set_cell(prev_table.cell(0, c), text, bold=True)
        previous = {p.slot: p for p in report.previous_findings}
        for r, slot in enumerate(range(1, 5), start=1):
            pf = previous.get(slot)
            if pf:
                _docx_set_cell(prev_table.cell(r, 0), _fmt_date(report.guidance_date))
                _docx_set_cell(prev_table.cell(r, 1), pf.title)
                _docx_set_cell(prev_table.cell(r, 2), pf.content)
                _docx_set_cell(prev_table.cell(r, 3), pf.action_result)
                _docx_set_cell(prev_table.cell(r, 4), "이행" if pf.confirmed else "-")

        document.add_page_break()

        # 5. 현재 공정 내 현존하는 위험성 제거
        document.add_heading("5. 현재 공정 내 현존하는 위험성 제거", level=2)
        risk_table = _docx_add_table(document, 5, 3)
        for c, text in enumerate(("유해·위험요인/지적사진", "지적사항(개선대책)", "비고")):
            _docx_set_cell(risk_table.cell(0, c), text, bold=True)
        findings = {f.slot: f for f in report.findings}
        for r, slot in enumerate(range(1, 5), start=1):
            f = findings.get(slot)
            if f:
                _docx_set_cell(risk_table.cell(r, 0), f.title)
                content = f.content + (f"\n* {f.law_citation}" if f.law_citation else "")
                _docx_set_cell(risk_table.cell(r, 1), content)
                risk_text = f"가능성 {f.likelihood} / 중대성 {f.severity}\n위험성 {f.risk_score}" if f.risk_score is not None else "-"
                _docx_set_cell(risk_table.cell(r, 2), risk_text, align="center")

        document.add_page_break()

        # 6. 향후 진행공정에 대한 유해·위험 요인 파악 및 대책
        document.add_heading("6. 향후 진행공정에 대한 유해·위험 요인 파악 및 대책", level=2)
        process_table = _docx_add_table(document, 1 + max(len(report.process_entries), 1), 4)
        for c, text in enumerate(("진행공정", "유해·위험요인", "예방대책", "위험성수준")):
            _docx_set_cell(process_table.cell(0, c), text, bold=True)
        for r, entry in enumerate(report.process_entries, start=1):
            _docx_set_cell(process_table.cell(r, 0), entry.process_name)
            _docx_set_cell(process_table.cell(r, 1), entry.hazard_text)
            _docx_set_cell(process_table.cell(r, 2), entry.prevention_text)
            _docx_set_cell(process_table.cell(r, 3), entry.risk_level, align="center")

        document.add_page_break()

        # 7. 사업장 지원 사항 등 기타 사항
        document.add_heading("7. 사업장 지원 사항 등 기타 사항", level=2)
        education = report.safety_education
        materials = report.provided_materials
        education_text = f"○ 참석인원: {education.attendee_count or '-'}명" if education else "○ 참석인원: -"
        material_names = "\n".join(f"- {m.title}" for m in materials) or "-"

        support_table = _docx_add_table(document, 3, 3)
        for c, text in enumerate(("지원사항", "구체적 사항", "비고")):
            _docx_set_cell(support_table.cell(0, c), text, bold=True)
        _docx_set_cell(support_table.cell(1, 0), "○ 교육", bold=True)
        _docx_set_cell(support_table.cell(1, 1), education_text)
        _docx_set_cell(support_table.cell(2, 0), "○ 자료배포", bold=True)
        _docx_set_cell(support_table.cell(2, 1), material_names)

        document.add_paragraph()
        measurement_by_type = {m.instrument_type: m for m in report.measurements}
        measurement_table = _docx_add_table(document, 1 + len(MEASUREMENT_INSTRUMENTS), 2)
        _docx_set_cell(measurement_table.cell(0, 0), "장비명", bold=True)
        _docx_set_cell(measurement_table.cell(0, 1), "측정값", bold=True)
        for r, (name, unit) in enumerate(MEASUREMENT_INSTRUMENTS, start=1):
            m = measurement_by_type.get(name)
            value = f"{m.value} {unit}" if m and m.value else "-"
            _docx_set_cell(measurement_table.cell(r, 0), name)
            _docx_set_cell(measurement_table.cell(r, 1), value, align="center")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))

        report.docx_path = str(output_path)
        session.commit()

    return output_path
