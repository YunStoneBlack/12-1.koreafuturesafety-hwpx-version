"""Report(회차 보고서)를 DOCX(python-docx)로 렌더링하는 모듈.

주의: 이 파일은 아직 예전 7섹션 구조 그대로다 — Sub-phase 7에서 PDF(report_builder_pdf.py)는
실제 표준 서식(9섹션)에 맞춰 다시 짰지만, DOCX는 현재 미리보기/생성 흐름에서 빠져있어서
(워드 버튼이 UI에 연결 안 됨) 이번 라운드에서는 손대지 않았다. 나중에 DOCX를 다시 연결할
때 report_builder_pdf.py의 9섹션 구조를 그대로 옮겨오면 된다.

HWPX는 core/hwpx_exporter.py가 이 DOCX를 한글에서 열어 다른 이름으로 저장하는 방식으로
만들기 때문에, DOCX 레이아웃이 정확할수록 HWPX 품질도 좋아진다.
"""

from __future__ import annotations

from pathlib import Path

from core.constants import FIXED_HAZARD_FACTORS, MEASUREMENT_INSTRUMENTS, NOTIFICATION_METHODS
from core.db import SessionLocal
from core.models_db import Report
from core.report_builder_common import fmt_amount, fmt_date


def _docx_set_cell(cell, text: str, bold: bool = False, align: str = "left") -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    for i, line in enumerate((text or "-").split("\n")):
        if i > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold


def _docx_add_table(document, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    return table


def build_report_docx(report_id: int, output_path: str | Path) -> Path:
    """Report 하나를 (예전 7섹션 구조 그대로) DOCX로 생성한다."""
    from docx import Document
    from docx.shared import Pt

    with SessionLocal() as session:
        report = session.get(Report, report_id)
        if report is None:
            raise ValueError(f"Report {report_id}를 찾을 수 없습니다.")
        site = report.site

        document = Document()
        style = document.styles["Normal"]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(9)

        document.add_heading("건설재해예방전문지도기관 기술지도 결과보고서", level=1)

        # 1. 기술지도 대상사업장
        document.add_heading("1. 기술지도 대상사업장", level=2)
        t = _docx_add_table(document, 7, 5)
        rows_data = [
            ("현장", "현장명", site.name, "사업장관리번호\n(사업개시번호)", f"{site.site_mgmt_no}\n{site.biz_start_no}"),
            ("", "공사기간", f"{fmt_date(site.period_start)} ~ {fmt_date(site.period_end)}", "공사금액", fmt_amount(site.amount)),
            ("", "책임자", site.manager_name, "연락처(이메일)", f"{site.manager_phone}\n{site.manager_email}"),
            ("", "주소", site.address, "", ""),
            ("본사", "회사명", site.hq_company, "법인등록번호\n(사업자등록번호)", f"{site.corp_reg_no}\n{site.biz_reg_no}"),
            ("", "면허번호", site.license_no, "연락처", site.hq_phone),
            ("", "주소", site.hq_address, "", ""),
        ]
        for r, row in enumerate(rows_data):
            for c, value in enumerate(row):
                _docx_set_cell(t.cell(r, c), value, bold=(c in (0, 1, 3)))

        # 2. 기술지도 개요
        document.add_heading("2. 기술지도 개요", level=2)
        method_text = " / ".join(f"{'■' if m == report.notification_method else '□'} {m}" for m in NOTIFICATION_METHODS)
        implemented = "□ 이행 / □ 불이행"
        if report.prev_guidance_implemented is True:
            implemented = "■ 이행 / □ 불이행"
        elif report.prev_guidance_implemented is False:
            implemented = "□ 이행 / ■ 불이행"
        staff_name = report.assigned_staff.name if report.assigned_staff else "-"
        staff_phone = report.assigned_staff.phone if report.assigned_staff else "-"

        t2 = _docx_add_table(document, 6, 4)
        rows2 = [
            ("지도기관명", "(주)한국미래안전", "기술지도실시일", fmt_date(report.guidance_date)),
            ("구분", "■ 건설 / □ 전기·정보통신", "공정률", f"({report.progress_rate or 0})%"),
            ("횟수", f"({report.visit_no})회차 / 총({site.total_guidance_count or '-'})회", "담당 요원", staff_name),
            ("이전 기술지도\n이행여부", implemented, "연락처", staff_phone),
            ("현장책임자 등\n통보방법", method_text, "", ""),
            ("기타 특이사항", report.special_note, "", ""),
        ]
        for r, row in enumerate(rows2):
            for c, value in enumerate(row):
                _docx_set_cell(t2.cell(r, c), value, bold=(c in (0, 2)))

        document.add_page_break()

        # 3. 현장 전경사진 + 12대 기인물 (예전 데이터 — OverviewPhoto가 비어있으면 빈 칸으로 나옴)
        document.add_heading("3. 현장 전경사진", level=2)
        overview_table = _docx_add_table(document, 2, 3)
        _docx_set_cell(overview_table.cell(0, 0), "위치", bold=True)
        _docx_set_cell(overview_table.cell(0, 1), "전경 1", bold=True)
        _docx_set_cell(overview_table.cell(0, 2), "전경 2", bold=True)
        _docx_set_cell(overview_table.cell(1, 0), "사진", bold=True)
        photos = {ph.slot: ph.photo_path for ph in report.overview_photos}
        for slot, col in ((1, 1), (2, 2)):
            path = photos.get(slot, "")
            cell = overview_table.cell(1, col)
            cell.text = ""
            if path and Path(path).exists():
                try:
                    cell.paragraphs[0].add_run().add_picture(path, width=Pt(120))
                except Exception:
                    _docx_set_cell(cell, "(이미지를 불러올 수 없음)")
            else:
                _docx_set_cell(cell, "(사진 없음)")

        document.add_paragraph()
        document.add_paragraph("건설현장 사망사고 다발 기인물 안전조치").runs[0].bold = True
        checked = set(report.hazard_factor_checks or [])
        hazard_table = _docx_add_table(document, 1 + len(FIXED_HAZARD_FACTORS), 3)
        for c, text in enumerate(("구분", "기인물", "안전조치")):
            _docx_set_cell(hazard_table.cell(0, c), text, bold=True)
        for i, (number, name, action) in enumerate(FIXED_HAZARD_FACTORS, start=1):
            mark = "■" if number in checked else "□"
            _docx_set_cell(hazard_table.cell(i, 0), mark, align="center")
            _docx_set_cell(hazard_table.cell(i, 1), f"{number}. {name}")
            _docx_set_cell(hazard_table.cell(i, 2), action)

        document.add_page_break()

        # 4. 이전 기술지도 사항 이행여부
        document.add_heading("4. 이전 기술지도 사항 이행여부", level=2)
        prev_table = _docx_add_table(document, 5, 5)
        for c, text in enumerate(("지도일(확인일)", "유해·위험요인/지적사진", "지적사항", "조치내용", "이행결과")):
            _docx_set_cell(prev_table.cell(0, c), text, bold=True)
        previous = {pf.slot: pf for pf in report.previous_findings}
        for r, slot in enumerate(range(1, 5), start=1):
            pf = previous.get(slot)
            if pf:
                _docx_set_cell(prev_table.cell(r, 0), fmt_date(report.guidance_date))
                _docx_set_cell(prev_table.cell(r, 1), pf.title)
                _docx_set_cell(prev_table.cell(r, 2), pf.content)
                _docx_set_cell(prev_table.cell(r, 3), pf.action_result)
                _docx_set_cell(prev_table.cell(r, 4), "이행" if pf.confirmed else "-")

        document.add_page_break()

        # 5. 현재 공정 내 현존하는 위험성 제거
        document.add_heading("5. 현재 공정 내 현존하는 위험성 제거", level=2)
        risk_table = _docx_add_table(document, 5, 3)
        for c, text in enumerate(("유해·위험요인/지적사진", "지적사항(개선대책)", "비고")):
            _docx_set_cell(risk_table.cell(0, c), text, bold=True)
        findings = {f.slot: f for f in report.findings}
        for r, slot in enumerate(range(1, 5), start=1):
            f = findings.get(slot)
            if f:
                _docx_set_cell(risk_table.cell(r, 0), f.title)
                content = f.content + (f"\n* {f.law_citation}" if f.law_citation else "")
                _docx_set_cell(risk_table.cell(r, 1), content)
                risk_text = f"가능성 {f.likelihood} / 중대성 {f.severity}\n위험성 {f.risk_score}" if f.risk_score is not None else "-"
                _docx_set_cell(risk_table.cell(r, 2), risk_text, align="center")

        document.add_page_break()

        # 6. 향후 진행공정에 대한 유해·위험 요인 파악 및 대책
        document.add_heading("6. 향후 진행공정에 대한 유해·위험 요인 파악 및 대책", level=2)
        process_table = _docx_add_table(document, 1 + max(len(report.process_entries), 1), 4)
        for c, text in enumerate(("진행공정", "유해·위험요인", "예방대책", "위험성수준")):
            _docx_set_cell(process_table.cell(0, c), text, bold=True)
        for r, entry in enumerate(report.process_entries, start=1):
            _docx_set_cell(process_table.cell(r, 0), entry.process_name)
            _docx_set_cell(process_table.cell(r, 1), entry.hazard_text)
            _docx_set_cell(process_table.cell(r, 2), entry.prevention_text)
            _docx_set_cell(process_table.cell(r, 3), entry.risk_level, align="center")

        document.add_page_break()

        # 7. 사업장 지원 사항 등 기타 사항
        document.add_heading("7. 사업장 지원 사항 등 기타 사항", level=2)
        education = report.safety_education
        materials = report.provided_materials
        education_text = f"○ 참석인원: {education.attendee_count or '-'}명" if education else "○ 참석인원: -"
        material_names = "\n".join(f"- {m.title}" for m in materials) or "-"

        support_table = _docx_add_table(document, 3, 3)
        for c, text in enumerate(("지원사항", "구체적 사항", "비고")):
            _docx_set_cell(support_table.cell(0, c), text, bold=True)
        _docx_set_cell(support_table.cell(1, 0), "○ 교육", bold=True)
        _docx_set_cell(support_table.cell(1, 1), education_text)
        _docx_set_cell(support_table.cell(2, 0), "○ 자료배포", bold=True)
        _docx_set_cell(support_table.cell(2, 1), material_names)

        document.add_paragraph()
        measurement_by_type = {m.instrument_type: m for m in report.measurements}
        measurement_table = _docx_add_table(document, 1 + len(MEASUREMENT_INSTRUMENTS), 2)
        _docx_set_cell(measurement_table.cell(0, 0), "장비명", bold=True)
        _docx_set_cell(measurement_table.cell(0, 1), "측정값", bold=True)
        for r, (name, unit) in enumerate(MEASUREMENT_INSTRUMENTS, start=1):
            m = measurement_by_type.get(name)
            value = f"{m.value} {unit}" if m and m.value else "-"
            _docx_set_cell(measurement_table.cell(r, 0), name)
            _docx_set_cell(measurement_table.cell(r, 1), value, align="center")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))

        report.docx_path = str(output_path)
        session.commit()

    return output_path
