"""결함조사만(2026-10-21 사용자 — 보고서 자동화 왼쪽 메뉴 "결함조사", 시특법과 설정 사이).
아직 보고서 틀이 없는 시설물(교량·수리 등)도 현장 결함조사는 쓰도록 — 시설물 이름·구분·지상/지하 층수·추가 탭(외부·정면·배면…)만 넣어
숨김 시설물(sitok_facility.survey_only, 시설물 목록·제출 현황·달력에 안 나옴) + 회차 하나를 만들고 현장 조사 화면(sitok-field.html)을 그대로 씀.
조사가 끝나면 결함현황표(엑셀)·사진첩(PDF — album_make: 층마다 결함 현황표 + 결함사진표) 내려받기. 전경·상태양호 등은 없음.

- `GET    /sitok/survey-only` — 목록
- `POST   /sitok/survey-only` {name, kind, floors_above, floors_below, roof, units} — 만들기
- `DELETE /sitok/survey-only/{facility_id}` {password} — 지우기(회사 삭제 비밀번호)
- `GET    /sitok/survey-only/{facility_id}/table` — 결함현황표 엑셀
- `GET    /sitok/survey-only/{facility_id}/album` — 사진첩 PDF
"""
from __future__ import annotations

import datetime
import io
from pathlib import Path

from fastapi import APIRouter, Body, Depends, HTTPException, status
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from core.models_web import SitokDefect, SitokFacility, SitokReport, User
from server.api.deps import get_current_user, get_db
from server.sitok import steps

router = APIRouter(prefix="/sitok", tags=["sitok"])
KINDS = ("건축물", "교량", "터널", "옹벽", "수리", "기타")


class SurveyIn(BaseModel):
    name: str
    kind: str = "건축물"
    floors_above: int | None = None
    floors_below: int | None = None
    roof: bool = True
    units: list[str] | None = None  # 추가 탭(외부·정면·배면…)


def _require(db: Session, user: User, facility_id: int) -> tuple[SitokFacility, SitokReport]:
    f = db.get(SitokFacility, facility_id)
    if f is None or f.company_id != user.company_id or not f.survey_only:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "결함조사를 찾을 수 없습니다.")
    r = db.query(SitokReport).filter(SitokReport.facility_id == f.id).order_by(SitokReport.id).first()
    if r is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "결함조사 회차가 없습니다.")
    return f, r


def _out(db: Session, f: SitokFacility, r: SitokReport | None) -> dict:
    from server.api.routers.sitok_defects import facility_floors

    rows = db.query(SitokDefect).filter(SitokDefect.report_id == r.id).all() if r is not None else []
    return {"id": f.id, "report_id": r.id if r is not None else None, "name": f.name, "kind": f.kind or "", "units": facility_floors(f),
            "defects": len(rows), "photos": sum(1 for d in rows if d.photo and Path(d.photo).exists()),
            "checked": sum(1 for d in rows if d.check), "created_by": f.created_by or "",
            "created_at": f.created_at.strftime("%Y-%m-%d") if getattr(f, "created_at", None) else "",
            "done": bool(steps.get(r, "field")) if r is not None else False}


@router.get("/survey-only")
def list_surveys(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    fs = db.query(SitokFacility).filter(SitokFacility.company_id == user.company_id, SitokFacility.survey_only.is_(True)).order_by(SitokFacility.id.desc()).all()
    out = []
    for f in fs:
        r = db.query(SitokReport).filter(SitokReport.facility_id == f.id).order_by(SitokReport.id).first()
        out.append(_out(db, f, r))
    return {"surveys": out, "kinds": KINDS}


@router.post("/survey-only")
def create_survey(body: SurveyIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    name = (body.name or "").strip()
    if not name:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "시설물 이름을 넣으세요.")
    kind = body.kind if body.kind in KINDS else "기타"
    above, below = max(int(body.floors_above or 0), 0), max(int(body.floors_below or 0), 0)
    units = [u.strip() for u in (body.units or []) if u and u.strip()]
    if not above and not below and not units and kind != "옹벽":
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "층수(지상·지하)나 탭(외부·정면 등)을 하나 이상 넣으세요.")
    led: dict = {"units_extra": list(dict.fromkeys(units))}
    if not body.roof and (above or below):
        led["units_hidden"] = ["옥상층"]
    f = SitokFacility(company_id=user.company_id, name=name, kind=kind, floors_above=above or None, floors_below=below or None,
                      survey_only=True, legal=False, template="2종", contract_type="1회성", sector="민간", created_by=user.display_name or "", ledger=led)
    db.add(f)
    db.flush()
    today = datetime.date.today()
    r = SitokReport(company_id=user.company_id, facility_id=f.id, year=today.year, half="상반기" if today.month <= 6 else "하반기", kind="정기",
                    period_start=today, source_note="결함조사만")
    db.add(r)
    db.commit()
    return _out(db, f, r)


@router.delete("/survey-only/{facility_id}")
def delete_survey(facility_id: int, password: str = Body("", embed=True), user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    from server.api.routers.sitok import _check_delete_password, _delete_facility

    f, _ = _require(db, user, facility_id)
    _check_delete_password(user, password)
    _delete_facility(db, f)
    db.commit()
    return {"ok": True}


STATUS_HEAD = ["층·탭", "번호", "구분", "부재", "결함유형", "개수", "폭/길이(mm/m)", "길이(m)", "물량(m/㎡)", "면적률(%)", "결함원인", "진행유무", "비고", "사진번호"]
STATUS_KEYS = ["번호", "구분", "부재", "결함유형", "개수", "폭", "길이", "물량", "면적률", "결함원인", "진행유무", "비고", "사진번호"]


@router.get("/survey-only/{facility_id}/table")
def download_table(facility_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """결함현황표 엑셀 — 층·탭마다 시트(사진첩 결함 현황표와 같은 칸) + 모두 한 시트."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from server.api.routers.sitok_defects import album_floors

    f, r = _require(db, user, facility_id)
    floors = album_floors(db, r.id)
    if not floors:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "아직 결함이 없습니다.")
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="888888")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, title: str, groups: list) -> None:
        ws.append([f"{f.name} 결함 현황표" + (f" — {title}" if title else "")])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(STATUS_HEAD)
        for c in ws[2]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8EDF5")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = box
        for fl, rows in groups:
            for row in rows:
                ws.append([fl] + [row.get(k, "") for k in STATUS_KEYS])
                for c in ws[ws.max_row]:
                    c.border = box
                    c.alignment = Alignment(vertical="center", wrap_text=True)
        for col, w in zip("ABCDEFGHIJKLMN", (10, 6, 9, 10, 12, 6, 10, 8, 9, 8, 24, 9, 7, 8)):
            ws.column_dimensions[col].width = w

    first = wb.active
    first.title = "전체"
    sheet(first, "", floors)
    for fl, rows in floors:
        sheet(wb.create_sheet(str(fl)[:31].replace("/", "·")), fl, [(fl, rows)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from urllib.parse import quote

    fname = quote(f"{f.name}_결함현황표.xlsx")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fname}"})


@router.get("/survey-only/{facility_id}/album")
def download_album(facility_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """사진첩 PDF — 층·탭마다 결함 현황표 쪽 + 결함사진표(한 쪽 6장)."""
    from server.api.routers.sitok import _facility_dir
    from server.api.routers.sitok_defects import album_floors
    from server.sitok import album_make

    f, r = _require(db, user, facility_id)
    floors = album_floors(db, r.id)
    if not floors:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "아직 결함이 없습니다.")
    import pymupdf

    from server.contract_docs import pdf_shrink

    dest = _facility_dir(f) / f"{f.name}_결함사진첩.pdf"
    dest.parent.mkdir(parents=True, exist_ok=True)
    raw = dest.with_name("_사진첩_원본.pdf")
    album_make.make(floors, raw)
    doc = pymupdf.open(raw)  # 따로 내려받는 책이라 여기서 줄임 — 맑은 고딕 통째(약 13MB)를 쓴 글자만, 폰 원본 사진은 150dpi(보고서 합본은 합칠 때 줄임)
    pdf_shrink.shrink_doc(doc)
    doc.subset_fonts()
    doc.save(dest, garbage=4, deflate=True)
    doc.close()
    raw.unlink(missing_ok=True)
    return FileResponse(dest, media_type="application/pdf", filename=dest.name, content_disposition_type="inline")
