"""실적(2026-10-14 사용자 — "실적은 내가 보기 위함, 계약 기준 얼마인지 파악해야 영업 공격성을 조절, 전년 동기 대비") — 서류 자동화 "실적" 탭.
시특법 계약(sitok_contract) + 산안법 용역계약(service_contract)에서 자동, 직접 입력 없음(프로그램 밖 업무는 이번엔 안 넣음).

- `GET /performance?year=2026` — {year, years, rows[계약 1건 = 1줄: law·kind·no·title·client·amount(그 해 몫)·date·start·end·halves], summary{this, prev(전년 같은 날짜까지), prev_full}, monthly{this[12], prev[12], law{시특법·산안법: {this, prev}}}}.
  금액 몫(portion): 시특법 민간 = 계약기간 안 반기마다 반기 금액(회차 줄이 있으면 그 금액) — 그 해에 드는 반기만 / 시특법 관급·산안법 = 계약일(없으면 착수일)이 든 해에 계약금액 전부.
- `GET /performance/xlsx?year=` — 같은 표 엑셀 한 장(+ 합계).
"""
from __future__ import annotations

import datetime
import io

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from core.models_web import ServiceContract, SitokContract, SitokFacility, User
from server.api.deps import get_current_user, get_db

router = APIRouter(prefix="/performance", tags=["performance"])
HALVES = ("상반기", "하반기")


def _halves_between(start: datetime.date, end: datetime.date) -> list[tuple[int, str, datetime.date]]:
    """계약기간에 드는 반기들 [(연도, 반기, 반기 시작일)] — 시작 반기부터 끝 반기까지."""
    out = []
    y, h = start.year, (0 if start.month <= 6 else 1)
    ey, eh = end.year, (0 if end.month <= 6 else 1)
    while (y, h) <= (ey, eh):
        out.append((y, HALVES[h], datetime.date(y, 1 if h == 0 else 7, 1)))
        y, h = (y, 1) if h == 0 else (y + 1, 0)
    return out


def _sitok_rows(db: Session, company_id: int) -> list[dict]:
    rows = []
    for c in db.query(SitokContract).filter(SitokContract.company_id == company_id).all():
        f = db.get(SitokFacility, c.facility_id)
        if f is None:
            continue
        kind = "시특법 " + ("학교" if (f.template or "").startswith("3종학교") else "건축") + f"·{c.sector or '민간'}"
        base = c.contract_date or c.start_date
        portions = []
        if (c.sector or "민간") == "민간" and c.start_date and c.end_date:
            by_round = {(r.get("year"), r.get("half")): r.get("amount") for r in (c.rounds or []) if r.get("year") and r.get("half")}
            for y, h, d in _halves_between(c.start_date, c.end_date):
                amt = by_round.get((y, h), c.amount) or 0
                portions.append({"date": d, "amount": int(amt), "label": f"{y} {h}"})
        elif c.amount and base:
            portions.append({"date": base, "amount": int(c.amount), "label": c.title or ""})
        rows.append({"law": "시특법", "kind": kind, "no": f.mgmt_no or "", "title": c.title or f"{f.name} 점검", "client": f.owner_name or "", "facility": f.name,
                     "date": base, "start": c.start_date, "end": c.end_date, "portions": portions, "link": f"sitok-facility.html?id={f.id}"})
    return rows


def _sanan_rows(db: Session, company_id: int) -> list[dict]:
    rows = []
    for c in db.query(ServiceContract).filter(ServiceContract.company_id == company_id).all():
        base = c.contract_date or c.start_date
        portions = [{"date": base, "amount": int(c.amount), "label": c.title or ""}] if c.amount and base else []
        rows.append({"law": "산안법", "kind": "기술지도·관급", "no": c.management_no or "", "title": c.title or "", "client": c.client or "", "facility": "",
                     "date": base, "start": c.start_date, "end": c.end_date, "portions": portions, "link": f"docs.html?contract={c.id}"})
    return rows


def _iso(d) -> str:
    return d.isoformat() if d else ""


def _build(db: Session, company_id: int, year: int) -> dict:
    today = datetime.date.today()
    all_rows = _sitok_rows(db, company_id) + _sanan_rows(db, company_id)
    years = sorted({p["date"].year for r in all_rows for p in r["portions"]} | {today.year}, reverse=True)
    rows = []
    for r in all_rows:
        mine = [p for p in r["portions"] if p["date"].year == year]
        if not mine:
            continue
        rows.append({**{k: v for k, v in r.items() if k not in ("portions", "date", "start", "end")}, "date": _iso(r["date"]), "start": _iso(r["start"]), "end": _iso(r["end"]),
                     "amount": sum(p["amount"] for p in mine), "total": sum(p["amount"] for p in r["portions"]),
                     "halves": [{"label": p["label"], "amount": p["amount"], "date": _iso(p["date"])} for p in mine]})
    rows.sort(key=lambda x: (x["date"] or "9999", x["no"]))

    def total(y: int, until: datetime.date | None = None) -> dict:
        out = {"total": 0, "sitok": 0, "sanan": 0, "count": 0}
        seen = set()
        for r in all_rows:
            for p in r["portions"]:
                if p["date"].year != y or (until and p["date"] > until):
                    continue
                out["total"] += p["amount"]
                out["sitok" if r["law"] == "시특법" else "sanan"] += p["amount"]
                seen.add(id(r))
        out["count"] = len(seen)
        return out

    cutoff = None
    if year == today.year:
        try:
            cutoff = today.replace(year=year - 1)
        except ValueError:  # 2월 29일
            cutoff = today.replace(year=year - 1, day=28)
    # 월별 — 합계(this·prev) + 구분별(law: 시특법·산안법, 화면 [전체·시특법·산안법] 버튼이 고름, 10/15 민재형 "그래프가 안 바뀜")
    monthly = {"this": [0] * 12, "prev": [0] * 12, "law": {k: {"this": [0] * 12, "prev": [0] * 12} for k in ("시특법", "산안법")}}
    for r in all_rows:
        for p in r["portions"]:
            key = "this" if p["date"].year == year else "prev" if p["date"].year == year - 1 else None
            if key:
                monthly[key][p["date"].month - 1] += p["amount"]
                monthly["law"][r["law"]][key][p["date"].month - 1] += p["amount"]
    return {"year": year, "years": years, "today": today.isoformat(), "rows": rows,
            "summary": {"this": total(year), "prev": total(year - 1, cutoff), "prev_full": total(year - 1), "cutoff": _iso(cutoff)}, "monthly": monthly}


@router.get("")
def get_performance(year: int | None = None, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return _build(db, user.company_id, year or datetime.date.today().year)


@router.get("/xlsx")
def performance_xlsx(year: int | None = None, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    y = year or datetime.date.today().year
    d = _build(db, user.company_id, y)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{y}년 실적"
    ws.append([f"{y}년 실적(계약 기준, VAT 포함)", "", "", "", "", "", "", f"만든 날 {d['today']}"])
    ws["A1"].font = Font(bold=True, size=13)
    head = ["번호", "구분", "용역명", "발주처·관리주체", f"{y}년 금액", "계약 총액", "계약일", "기간", "반기 내역"]
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF8")
    for r in d["rows"]:
        ws.append([r["no"], r["kind"], r["title"], r["client"], r["amount"], r["total"], r["date"], f"{r['start']} ~ {r['end']}", " / ".join(f"{h['label']} {h['amount']:,}" for h in r["halves"])])
    s = d["summary"]
    ws.append([])
    ws.append(["합계", "", "", "", s["this"]["total"], "", "", f"시특법 {s['this']['sitok']:,} · 산안법 {s['this']['sanan']:,}", f"전년 같은 기간 {s['prev']['total']:,}"])
    ws[ws.max_row][0].font = Font(bold=True)
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for c in (row[4], row[5]):
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
    for col, w in zip("ABCDEFGHI", (9, 16, 40, 24, 14, 14, 12, 24, 40)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from urllib.parse import quote

    name = quote(f"실적_{y}.xlsx")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{name}"})
