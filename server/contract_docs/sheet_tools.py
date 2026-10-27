"""양식 시트 손질 — openpyxl로 연 회사 양식(착수계·완수계)의 그림·도장·찌꺼기를 다룬다.

양식 속 그림(10/6 확인):
- 둥근 대표이사 도장(237×238 그림) — 시트마다 인쇄 범위 밖(오른쪽)에 놓여 있다. [도장 넣기]면 "대표 이 인 숙 (인)" 위나
  "원본대조필" 칸 오른쪽에 새로 놓고, 양식에 있던 것은 늘 지운다(빼기면 인쇄 뒤 실제 도장).
- 원본대조필 칸(209×64), 감독경유 칸(475×280), 갑지 네모 직인·로고 — 그대로 둔다.
- 증명서·자격증 사진 — 시트에서 가장 큰 그림. 올린 파일로 바꾼다(같은 자리·같은 비율 — 비율이 다르면 흰 여백을 붙임).

좌표: 열 너비(글자 수) → 픽셀 ≈ 너비×7+5, 행 높이(pt) → 픽셀 = pt×96/72. 도장처럼 몇 픽셀 어긋나도 되는 곳에만 쓴다.
"""
from __future__ import annotations

import io
from copy import copy
from pathlib import Path

from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PILImage
from PIL import ImageOps

SEAL_PX = (237, 238)
SEAL_SIZE = 66  # 인쇄 때 도장 크기(픽셀, 약 17mm)
MAX_DOC_PX = 1600  # 증명서 그림 긴 변 — 파일이 너무 커지지 않게(A4 한 장에 충분)


def col_px(ws: Worksheet, col: int) -> float:
    """열(1부터) 너비 픽셀."""
    dim = ws.column_dimensions.get(get_column_letter(col))
    width = dim.width if dim is not None and dim.width else (ws.sheet_format.defaultColWidth or 8.43)
    if dim is not None and dim.hidden:
        return 0
    return width * 7 + 5


def row_px(ws: Worksheet, row: int) -> float:
    dim = ws.row_dimensions.get(row)
    pt = dim.height if dim is not None and dim.height else (ws.sheet_format.defaultRowHeight or 15)
    if dim is not None and dim.hidden:
        return 0
    return pt * 96 / 72


def col_x(ws: Worksheet, col: int) -> float:
    """열(1부터) 왼쪽 끝 x 픽셀."""
    return sum(col_px(ws, c) for c in range(1, col))


def row_y(ws: Worksheet, row: int) -> float:
    return sum(row_px(ws, r) for r in range(1, row))


def _marker(ws: Worksheet, x: float, y: float) -> AnchorMarker:
    col, left = 1, x
    while left >= col_px(ws, col) and col < 200:
        left -= col_px(ws, col)
        col += 1
    row, top = 1, y
    while top >= row_px(ws, row) and row < 2000:
        top -= row_px(ws, row)
        row += 1
    return AnchorMarker(col=col - 1, colOff=pixels_to_EMU(max(0, left)), row=row - 1, rowOff=pixels_to_EMU(max(0, top)))


def print_bounds(ws: Worksheet) -> tuple[int, int]:
    """인쇄 범위의 (마지막 열, 마지막 행) — 1부터."""
    area = ws.print_area
    if not area:
        return ws.max_column, ws.max_row
    ref = (area[0] if isinstance(area, (list, tuple)) else area).split("!")[-1].replace("$", "")
    _, _, max_col, max_row = range_boundaries(ref)
    return max_col, max_row


def fix_dot_numbers(wb) -> None:
    """번호 칸 서식 "."(엑셀에선 "1."로 보이지만 LibreOffice PDF는 "."만 — 착수계·완수계 PDF 줄 번호가 빠져 있었음, 10/20) → 0"."."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.number_format == "." and isinstance(c.value, (int, float)):
                    c.number_format = '0"."'


def clear_outside_print(ws: Worksheet) -> None:
    """인쇄 범위 오른쪽 칸의 값 지우기 — 양식을 다른 현장에서 베껴 오며 남은 찌꺼기(화현초·초가팔2리 등, 인쇄엔 안 나옴)."""
    max_col, _ = print_bounds(ws)
    for row in ws.iter_rows(min_col=max_col + 1):
        for c in row:
            if c.value is not None and not isinstance(c, MergedCell):
                c.value = None


def img_bytes(img: XLImage) -> bytes:
    """그림 바이트. openpyxl은 한 번 읽으면 파일을 닫아 두 번째(저장 때 포함)에 실패하므로 읽은 뒤 새 BytesIO로 되돌려 둔다."""
    data = getattr(img, "_kfsc_bytes", None)
    if data is None:
        data = img._data()
        img._kfsc_bytes = data
    img.ref = io.BytesIO(data)
    return data


def _px(img: XLImage) -> tuple[int, int]:
    try:
        with PILImage.open(io.BytesIO(img_bytes(img))) as p:
            return p.size
    except Exception:  # noqa: BLE001
        return (0, 0)


def is_seal(img: XLImage) -> bool:
    return _px(img) == SEAL_PX


def seal_bytes(wb) -> bytes | None:
    """양식에 들어 있는 둥근 대표이사 도장 그림(첫 번째)."""
    for ws in wb.worksheets:
        for img in ws._images:
            if is_seal(img):
                return img_bytes(img)
    return None


def remove_seals(ws: Worksheet) -> None:
    ws._images = [img for img in ws._images if not is_seal(img)]


def _anchor_box(ws: Worksheet, img: XLImage) -> tuple[float, float, float, float]:
    """그림의 (x, y, 너비, 높이) 픽셀 — 두 칸 기준 위치만(양식 그림은 전부 두 칸 기준)."""
    a = img.anchor
    fx = col_x(ws, a._from.col + 1) + a._from.colOff / 9525
    fy = row_y(ws, a._from.row + 1) + a._from.rowOff / 9525
    to = getattr(a, "to", None)
    if to is None:
        return fx, fy, 0, 0
    tx = col_x(ws, to.col + 1) + to.colOff / 9525
    ty = row_y(ws, to.row + 1) + to.rowOff / 9525
    return fx, fy, tx - fx, ty - fy


def main_image(ws: Worksheet) -> XLImage | None:
    """증명서·자격증 사진 — 인쇄 범위 안에서 가장 큰 그림(도장 빼고)."""
    max_col, _ = print_bounds(ws)
    best, area = None, 0.0
    for img in ws._images:
        if is_seal(img) or img.anchor._from.col >= max_col:
            continue
        _, _, w, h = _anchor_box(ws, img)
        if w * h > area:
            best, area = img, w * h
    return best


def find_image(ws: Worksheet, px: tuple[int, int]) -> XLImage | None:
    for img in ws._images:
        if _px(img) == px:
            return img
    return None


MARGIN_MM = 3  # 서류 그림 칸 안 여백(사방, 10/27 사용자 "원본 비율 최대한 유지하면서 틀에 맞게, 여백 3미리")


def _upright(src: str | Path | bytes) -> bytes:
    """올린 그림 → 바로 세운(폰 사진 회전 정보 — 카메라로 찍은 성과품 표지 등) RGB JPEG, 긴 변 MAX_DOC_PX."""
    with PILImage.open(io.BytesIO(src) if isinstance(src, bytes) else str(src)) as p:
        p = ImageOps.exif_transpose(p).convert("RGB")
        p.thumbnail((MAX_DOC_PX, MAX_DOC_PX))
        buf = io.BytesIO()
        p.save(buf, "JPEG", quality=88)
        return buf.getvalue()


def _pad(content: bytes, box_w: float, box_h: float, margin: float = 0) -> bytes:
    """content(JPEG)를 칸(box_w × box_h — 단위는 margin과 같음) 비율의 흰 바탕 가운데에, 사방 margin을 남기고 원본 비율 그대로 가장 크게 → JPEG.
    칸에 그대로 늘려 넣어도(두 칸 기준 그림) 그림이 안 찌그러짐."""
    with PILImage.open(io.BytesIO(content)) as p:
        p = p.convert("RGB")
        w, h = p.size
        if box_w <= 0 or box_h <= 0 or not w or not h:
            canvas = p
        else:
            iw, ih = max(box_w - 2 * margin, box_w * 0.5), max(box_h - 2 * margin, box_h * 0.5)
            s = min(iw / w, ih / h)  # 칸 단위 / 그림 px
            cw, ch = max(round(box_w / s), w), max(round(box_h / s), h)
            canvas = PILImage.new("RGB", (cw, ch), "white")
            canvas.paste(p, ((cw - w) // 2, (ch - h) // 2))
            if max(canvas.size) > MAX_DOC_PX * 1.3:  # 바탕이 너무 크면 줄임(파일 크기)
                canvas.thumbnail((round(MAX_DOC_PX * 1.3), round(MAX_DOC_PX * 1.3)))
        buf = io.BytesIO()
        canvas.save(buf, "JPEG", quality=88)
        return buf.getvalue()


def _frame_anchor(ws: Worksheet, img: XLImage):
    """그림이 들어 있는 테두리 틀(합친 칸 — 결과보고서·통장 사본 B4:M25, 4대보험 B4:M27)에 꼭 맞는 두 칸 기준 자리. 틀이 없으면 None.
    (10/27 — 견본 그림이 틀 위쪽 일부에만 걸려 있어 사진 윗부분이 제목 줄 뒤로 가려지고 틀 아래는 비던 것)"""
    a = img.anchor
    fr, fc = a._from.row + 1, a._from.col + 1
    best = None
    for m in ws.merged_cells.ranges:
        if m.min_row <= fr <= m.max_row and m.min_col <= fc <= m.max_col and m.max_row - m.min_row >= 4 and m.max_col - m.min_col >= 3:
            if best is None or (m.max_row - m.min_row) * (m.max_col - m.min_col) > (best.max_row - best.min_row) * (best.max_col - best.min_col):
                best = m
    if best is None:
        return None
    m = MARGIN_MM / 25.4 * 96  # 틀 안쪽으로 사방 3mm 들임 — 그림 흰 바탕이 틀 테두리 선을 덮지 않게(여백은 그림 밖)

    def end(last: int, size) -> tuple[int, float]:  # 끝 칸(1부터)에서 m만큼 안쪽 — 끝 칸이 좁으면 앞 칸으로
        k, back = last, m
        while k > 1 and size(k) <= back + 1:
            back -= size(k)
            k -= 1
        return k - 1, max(size(k) - back, 1)

    tc, tco = end(best.max_col, lambda c: col_px(ws, c))
    tr, tro = end(best.max_row, lambda r: row_px(ws, r))
    return TwoCellAnchor(_from=AnchorMarker(col=best.min_col - 1, row=best.min_row - 1, colOff=pixels_to_EMU(m), rowOff=pixels_to_EMU(m)),
                         to=AnchorMarker(col=tc, row=tr, colOff=pixels_to_EMU(tco), rowOff=pixels_to_EMU(tro)))


def replace_image(ws: Worksheet, old: XLImage, src: str | Path | bytes) -> None:
    """old 그림 자리에 src(파일 경로·바이트)를 넣는다 — 자리는 old가 든 테두리 틀(합친 칸) 전체, 틀이 없으면 old 그대로.
    원본은 그림에 붙여 둠(refit_pictures가 실제 칸 비율로 다시)."""
    frame = _frame_anchor(ws, old)
    if frame is not None:
        old.anchor = frame
    margin = 0 if frame is not None else MARGIN_MM  # 틀이 있으면 여백은 그림 밖(들인 자리), 없으면 그림 안 흰 바탕
    _, _, w, h = _anchor_box(ws, old)
    content = _upright(src)
    data = _pad(content, w, h, margin / 25.4 * 96)
    new = XLImage(io.BytesIO(data))
    new.anchor = old.anchor
    new._kfsc_src, new._kfsc_data, new._kfsc_margin = content, data, margin
    with PILImage.open(io.BytesIO(data)) as p:
        new._kfsc_px = p.size
    ws._images[ws._images.index(old)] = new


def refit_pictures(wb, xlsx: Path, margin_mm: float = MARGIN_MM) -> int:
    """replace_image로 넣은 서류 그림을 실제 칸 비율로 다시(10/27 — 엑셀 열 너비·행 높이 어림과 LibreOffice PDF의 실제 칸이 달라
    결과보고서 표지가 가로로 30%·4대보험 서류가 10~15% 눌리던 것): xlsx를 PDF로 한 번 바꿔 그림 칸(넣은 그림 px로 찾음)을 재고,
    원본 비율 그대로 사방 margin_mm 안에 가장 크게 → xlsx 다시 저장. 바꾼 그림 수(못 하면 0 — 그때는 어림 그대로)."""
    import tempfile

    import pymupdf

    from server.contract_docs import to_pdf

    marked = [(ws, img) for ws in wb.worksheets for img in ws._images if getattr(img, "_kfsc_src", None) is not None]
    if not marked:
        return 0
    try:
        with tempfile.TemporaryDirectory() as tmp:
            pdf = to_pdf.office_to_pdf(Path(xlsx), Path(tmp) / "fit.pdf")
            with pymupdf.open(pdf) as doc:
                rects = [((info[2], info[3]), r) for pg in doc for info in pg.get_images(full=True) for r in pg.get_image_rects(info[0])]
    except Exception:  # noqa: BLE001
        return 0
    used: set = set()
    swap: dict[bytes, bytes] = {}  # 엑셀 안 그림 파일 바이트 → 새 바이트(저장한 통합문서를 다시 저장하면 앞 그림 파일이 닫혀 실패 — zip 속 그림만 바꿈)
    for ws, img in marked:
        px0 = tuple(img._kfsc_px)
        ok = [i for i, (px, r) in enumerate(rects) if i not in used and r.width > 20 and r.height > 20]
        k = next((i for i in ok if rects[i][0] == px0), None)  # 넣은 그림 px 그대로
        if k is None:  # PDF 변환기가 큰 그림 해상도를 줄임(통장 사본 1856×1664 → 1747×1566) — 쪽 차례대로 가로세로 비율(1% 안)로
            k = next((i for i in ok if rects[i][0][1] and abs(rects[i][0][0] / rects[i][0][1] - px0[0] / px0[1]) < 0.01 * px0[0] / px0[1]), None)
        if k is None:
            continue
        used.add(k)
        r = rects[k][1]
        swap[img._kfsc_data] = _pad(img._kfsc_src, r.width, r.height, getattr(img, "_kfsc_margin", margin_mm) / 25.4 * 72)
    if not swap:
        return 0
    import shutil
    import zipfile

    tmp_x = Path(xlsx).with_suffix(".refit.xlsx")
    n = 0
    with zipfile.ZipFile(xlsx) as zin, zipfile.ZipFile(tmp_x, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.startswith("xl/media/") and data in swap:
                data = swap[data]
                n += 1
            zout.writestr(info, data)
    shutil.move(str(tmp_x), str(xlsx))
    return n


def drop_image(ws: Worksheet, img: XLImage | None) -> None:
    if img is not None and img in ws._images:
        ws._images.remove(img)


def add_image_copy(ws: Worksheet, img: XLImage) -> None:
    """다른 시트 그림을 같은 자리에 복사(시트 복사 때 — openpyxl copy_worksheet는 그림을 안 옮김)."""
    new = XLImage(io.BytesIO(img_bytes(img)))
    new.anchor = copy(img.anchor)
    ws._images.append(new)


def place_seal(ws: Worksheet, seal: bytes, cx: float, cy: float, size: int = SEAL_SIZE) -> None:
    """도장 그림을 (cx, cy) 픽셀이 가운데가 되게 놓는다."""
    img = XLImage(io.BytesIO(seal))
    img.anchor = OneCellAnchor(_from=_marker(ws, cx - size / 2, cy - size / 2),
                               ext=XDRPositiveSize2D(pixels_to_EMU(size), pixels_to_EMU(size)))
    ws._images.append(img)


def anchor_box(ws: Worksheet, img: XLImage) -> tuple[float, float, float, float]:
    """그림의 (x, y, 너비, 높이) 픽셀 — 한 칸 기준(크기 ext) 그림도."""
    x, y, w, h = _anchor_box(ws, img)
    ext = getattr(img.anchor, "ext", None)
    if not w and ext is not None:
        w, h = ext.width / 9525, ext.height / 9525
    return x, y, w, h


def place_picture(ws: Worksheet, data: bytes, box: tuple[float, float, float, float]) -> None:
    """그림(PNG 등 — 투명 그대로)을 칸(x, y, 너비, 높이 픽셀) 안에 비율대로 가운데 맞춰 놓는다(시특법 보안각서 서명, 10/12)."""
    x, y, w, h = box
    with PILImage.open(io.BytesIO(data)) as p:
        pw, ph = p.size
    k = min(w / pw, h / ph) if pw and ph else 1
    iw, ih = pw * k, ph * k
    img = XLImage(io.BytesIO(data))
    img.anchor = OneCellAnchor(_from=_marker(ws, x + (w - iw) / 2, y + (h - ih) / 2), ext=XDRPositiveSize2D(pixels_to_EMU(iw), pixels_to_EMU(ih)))
    ws._images.append(img)


def merged_right_x(ws: Worksheet, coord: str) -> float:
    """칸(병합이면 병합 범위)의 오른쪽 끝 x 픽셀."""
    cell = ws[coord]
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            return col_x(ws, rng.max_col + 1)
    return col_x(ws, cell.column + 1)


def seal_on_text_end(ws: Worksheet, coord: str, seal: bytes, back_px: float = 46) -> None:
    """오른쪽 정렬된 "대표 이 인 숙 (인)" 칸 — 글자 끝 "(인)" 위에 도장(끝에서 back_px 안쪽이 가운데)."""
    cell = ws[coord]
    cx = merged_right_x(ws, coord) - back_px
    cy = row_y(ws, cell.row) + row_px(ws, cell.row) / 2
    place_seal(ws, seal, cx, cy)


def seal_on_original_box(ws: Worksheet, seal: bytes) -> bool:
    """"원본대조필" 칸(209×64 그림)의 오른쪽 빈 칸 위에 도장. 칸이 없으면 False."""
    box = find_image(ws, (209, 64))
    if box is None:
        return False
    x, y, w, h = _anchor_box(ws, box)
    place_seal(ws, seal, x + w * 0.78, y + h / 2)
    return True
