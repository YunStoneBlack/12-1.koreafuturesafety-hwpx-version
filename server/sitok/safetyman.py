"""상태평가 결과 자료 = SafetyMan 출력과 똑같이(2026-10-15 민재형 "세이프티맨 출력물과 일치해야 하는데 디테일이 너무 떨어짐 — 글자 크기·두께·색깔·전경사진·
구조형식 뒤 한 칸… 100% 일치"). 방법: 형이 준 SafetyMan 출력 엑셀(숭의 본관 2026 상반기)을 양식으로 — 좁은 열(0.89)을 바둑판처럼 깔고 칸을 합쳐 그린
SafetyMan 짜임을 블록(줄 묶음)째 복사해 층 묶음·평가부재 수만큼 늘리고 우리 값을 넣은 뒤 LibreOffice로 PDF. (양식 그대로를 LibreOffice로 바꾸면
SafetyMan PDF와 쪽 수·모양이 같게 나옴 — 10/15 확인.)

양식 줄(블록) — data/templates/시특법_SafetyMan_양식.xlsx 첫 시트:
  1~30 제목·건물개요(전경사진 = 10~25행 AU~CE) / 31~34 평가결과 머리 / 35~37 층 묶음 한 줄(부재 머리·상태, 기울기 및 침하는 BX 세로 합침) /
  47~50 최종결과 / 51~52 입력자료 / 53~54 묶음 제목 / 55~56·57·59 콘크리트강도(평가부재 줄 반복) / 60~69 균열 / 70~71·72·73 중성화 /
  74~83 박리 / 84~93 박락·층분리 / 94~101 누수·백태 / 102~109 철근노출 / 110 묶음 사이 / 281~282·283·285 강재규격 / 286~293 용접 /
  294~297 볼트 / 298~305 강재부식도 / 306~309 내화피복 / 310~315 기울기 및 침하·메모.
SRC 층은 콘크리트 묶음 "라멘(RC)" + 같은 층 "N층(N층 ~ N층) 라멘(철골조)" 묶음(형 SafetyMan 입력 방식 — 집계가 steel_part 층으로 따로 평가, condition.aggregate).
"""
from __future__ import annotations

import copy
import io
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor

from server.contract_docs.build import TEMPLATE_DIR
from server.sitok import rating

TEMPLATE = TEMPLATE_DIR / "시특법_SafetyMan_양식.xlsx"
MEMBERS = rating.MEMBERS  # 기둥 내력벽 큰보 작은보 슬래브 테두리보 접합부
RES_COLS = ("T", "AA", "AH", "AO", "AV", "BC", "BJ")  # 평가결과 부재 칸, 종합 = BQ
DS_COLS = ("N", "X", "AH", "AR", "BB", "BL", "BV")  # 강도·중성화 평가부재 칸
CNT_COLS = ("AO", "AU", "BA", "BG", "BM", "BS", "BZ")  # 균열·박리… 개수 칸
STEEL_DS_COLS = {"기둥": "N", "큰보": "X", "작은보": "AH"}  # 강재규격 칸
STEEL_CNT_COLS = {"기둥": "AO", "큰보": "AU", "작은보": "BA"}  # 용접·볼트·부식 칸
# 기울기·부동침하 각변위 등급 → (각변위, 상태) — A·C는 SafetyMan 출력 그대로, B·D·E는 세부지침 표 6.35 글(SafetyMan 출력 못 봄)
DEFORM_TEXT = {"a": ("1/750 이하", "A : 예민한 기계 기초의 위험 침하 한계"), "b": ("1/500 이하", "B : 균열을 허용하지 않는 건물의 안전 한계"),
               "c": ("1/250 이하", "C : 구조물의 경사도 감지"), "d": ("1/150 이하", "D : 구조적 손상이 예상되는 한계"), "e": ("1/150 초과", "E : 구조적 손상 발생")}


def _floor_name(n: int) -> str:
    return f"{n}층"  # SafetyMan: 지하 = 음수("-1층")


class Book:
    """양식 시트에서 블록을 복사해 새 시트에 차례로 쌓기."""

    def __init__(self):
        self.wb = openpyxl.load_workbook(TEMPLATE)
        self.src = self.wb.worksheets[0]
        self.dst = self.wb.create_sheet("평가결과")
        for k, d in self.src.column_dimensions.items():
            nd = self.dst.column_dimensions[k]
            nd.width, nd.min, nd.max = d.width, d.min, d.max
        self.dst.sheet_format = copy.copy(self.src.sheet_format)
        self.dst.page_margins = copy.copy(self.src.page_margins)
        ps, nps = self.src.page_setup, self.dst.page_setup
        nps.orientation, nps.paperSize, nps.scale = ps.orientation, ps.paperSize, ps.scale
        self.dst.print_options = copy.copy(self.src.print_options)
        self.merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in self.src.merged_cells.ranges]
        self.row = 1

    def block(self, a: int, b: int, skip_merge=None) -> int:
        """양식 a~b행을 지금 줄에 복사(값·모양·높이·그 안의 합침) → 시작 줄."""
        start = self.row
        for r in range(a, b + 1):
            t = start + r - a
            self.dst.row_dimensions[t].height = self.src.row_dimensions[r].height
            for c in range(1, self.src.max_column + 1):
                s = self.src.cell(r, c)
                d = self.dst.cell(t, c)
                if s.has_style:
                    d._style = copy.copy(s._style)
                if not isinstance(s, MergedCell) and s.value is not None:
                    d.value = s.value
        for (r1, c1, r2, c2) in self.merges:
            if a <= r1 and r2 <= b and not (skip_merge and skip_merge(r1, c1, r2, c2)):
                self.dst.merge_cells(start_row=start + r1 - a, start_column=c1, end_row=start + r2 - a, end_column=c2)
        self.row = start + (b - a + 1)
        return start

    def put(self, col: str, row: int, value) -> None:
        self.dst[f"{col}{row}"].value = value

    def clear(self, cols, row: int) -> None:
        for c in cols:
            self.put(c, row, None)

    def photo(self, path: Path, r0: int = 9, c0: int = 46, r1: int = 25, c1: int = 83) -> None:
        """전경사진 — 건물개요 오른쪽 틀(양식과 같은 자리)에 맞게 가운데 잘라 넣음."""
        from PIL import Image, ImageOps

        img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
        box_w, box_h = (c1 - c0) * 11.2, (r1 - r0) * 18.0  # 열 0.89 ≈ 11.2px, 줄 13.5pt = 18px
        want = box_w / box_h
        w, h = img.size
        if w / h > want:
            nw = int(h * want)
            img = img.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
        else:
            nh = int(w / want)
            img = img.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
        img.thumbnail((1000, 1000))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=88)
        buf.seek(0)
        xi = XlImage(buf)
        xi.width, xi.height = int(box_w), int(box_h)
        xi.anchor = TwoCellAnchor(editAs="twoCell", _from=AnchorMarker(col=c0, row=r0), to=AnchorMarker(col=c1, row=r1))
        self.dst.add_image(xi)

    def save(self, dest: Path) -> Path:
        for ws in list(self.wb.worksheets):
            if ws is not self.dst:
                self.wb.remove(ws)
        self.wb.active = 0
        self.dst.print_area = f"A1:CF{self.row}"
        dest.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(dest)
        return dest


def _n(x) -> str:
    """SafetyMan 숫자 모양 — 24.0 → "24", 24.10 → "24.1", 글은 그대로."""
    try:
        v = float(x)
    except (TypeError, ValueError):
        return str(x if x is not None else "")
    return f"{v:.2f}".rstrip("0").rstrip(".")


def _groups(agg: dict, result: dict | None) -> list[dict]:
    """평가결과·입력자료 묶음 — 집계 층(입력)과 결과 층(점수)은 같은 차례(condition.aggregate). SRC의 강재 쪽(steel_part)은 맨 뒤 철골조 묶음."""
    res = list((result or {}).get("floors", []))
    pairs = [(fl, res[i] if i < len(res) else None) for i, fl in enumerate(agg["floors"])]
    order = sorted([p for p in pairs if not p[0].get("steel_part")], key=lambda p: -p[0]["floor"]) + [p for p in pairs if p[0].get("steel_part")]  # SafetyMan: 위층부터(12층 → 5층 → -1층)
    out = []
    for fl, r in order:
        steel = fl["material"] == "철골"  # 철골조 층 — SafetyMan "라멘(철골조)" 묶음(콘크리트 표 + 강재 표)
        mat = "철골조" if steel else "RC" if fl["material"] == "SRC" else fl["material"]
        out.append({"agg": fl, "res": r, "kind": "steel" if steel else "rc", "name": _floor_name(fl["floor"]), "range": fl.get("range_label", ""),
                    "structure": f"{fl['structure']}({mat})",
                    "steel": [(m["kind"], x) for m in fl["members"] for x in (((m.get("steel") or {}).get("SS") or {}).get("rows", []))]})
    return out


def _range_text(fl: dict) -> str:
    """"지하1층~지상3층" → "-1층 ~ 3층"(SafetyMan 모양)."""
    import re

    def one(t):
        t = t.strip()
        m = re.match(r"지하(\d+)층", t)
        if m:
            return f"-{m[1]}층"
        m = re.match(r"(?:지상)?(\d+)층", t)
        return f"{m[1]}층" if m else t

    parts = [p for p in re.split(r"\s*~\s*", fl or "") if p]
    return " ~ ".join(one(p) for p in parts) if parts else ""


def build(dest: Path, agg: dict, result: dict | None, building: dict, tests_view: dict, photo: Path | None = None, joints: dict | None = None) -> Path:
    """SafetyMan 모양 엑셀(dest .xlsx)."""
    B = Book()
    bi = building
    # ---- 건물개요
    B.block(1, 30)
    vals = {7: bi.get("name", ""), 8: bi.get("address", ""), 9: bi.get("completion", ""), 10: bi.get("use", ""), 11: bi.get("structure", ""),
            13: bi.get("owner", ""), 14: "", 15: bi.get("manager") or bi.get("owner", ""), 16: "",
            18: bi.get("land_area") or "-", 19: bi.get("building_area") or "-", 20: bi.get("total_area") or "-",
            21: bi.get("coverage") or "-", 22: bi.get("far") or "-", 24: f"{bi.get('above', '')} 층", 25: f"{bi.get('below', '')} 층"}
    for r, v in vals.items():
        B.put("N", r, v)
    B.put("N", 27, bi.get("kind_label") or "정밀점검")
    if bi.get("company"):  # 평가기관·연락처 — 안 주면 양식 값 그대로(SafetyMan 설정: (주)한국미래안전·031-542-0456)
        B.put("BC", 27, bi["company"])
    B.put("N", 28, bi.get("when", ""))
    B.dst["N28"].number_format = "@"
    if bi.get("phone"):
        B.put("BC", 28, bi["phone"])
    if photo and Path(photo).exists():
        B.photo(Path(photo))
    # ---- 평가결과
    B.block(31, 34)
    groups = _groups(agg, result)
    first = B.row
    for g in groups:
        r0 = B.block(35, 37, skip_merge=lambda r1, c1, r2, c2: c1 >= 76)  # BX 세로 칸은 끝에 한 번
        B.put("B", r0, f"{g['name']} \n({_range_text(g['range'])})\n{g['structure']}")
        B.clear(RES_COLS + ("BQ", "BX"), r0 + 1)
        B.put("BX", r0, None)
        scores, total = {}, None
        if g["res"]:
            scores = {m["kind"]: m.get("condition") for m in g["res"].get("members", []) if m["kind"] in MEMBERS}
            total = g["res"].get("condition")
        for k, col in zip(MEMBERS, RES_COLS):
            v = scores.get(k)
            B.put(col, r0 + 1, "-" if v is None else round(float(v), 2))
        B.put("BQ", r0 + 1, f"{total:.2f}({rating.grade_of(total)})" if total is not None else "-")
    last = B.row - 1
    if groups:
        B.dst.merge_cells(f"BX{first}:CE{last}")
        d = (result or {}).get("deform")
        B.put("BX", first, f"{d['score']:.2f}({d['grade']})" if d else "-")
    r = B.block(47, 50)
    if result and result.get("condition") and result.get("total"):
        B.put("N", r, f"상태평가: {result['condition']['score']:.2f}({result['condition']['grade']}등급)\n종합평가: {result['total']['score']:.2f}({result['total']['grade']}등급)")
    # ---- 입력자료
    B.block(51, 52)
    for i, g in enumerate(groups):
        fl = g["agg"]
        mem = {m["kind"]: m for m in fl["members"]}
        r = B.block(53, 54)
        B.put("B", r, f"{g['name']}[{_range_text(g['range']).replace(' ', '')}]{g['structure']}")
        _test_rows(B, mem, "DS", (55, 56), 57, 59, (229, 231), lambda x: f"{_n(x.get('est'))}, 0,\n0 / {_n(x.get('design'))}" + ("(D)" if x.get("damaged") else ""))
        _count_table(B, mem, (60, 69), "CR", None, area=True)
        _test_rows(B, mem, "CA", (70, 71), 72, 73, (242, 244), lambda x: f"{_n(x.get('depth'))}, 0, 0 \n/ {_n(x.get('cover'))}")
        _count_table(B, mem, (74, 83), "SD", "scaling", area=True)
        _count_table(B, mem, (84, 93), "SD", "spalling", area=True)
        _count_table(B, mem, (94, 101), "SD", "leak", area=False)
        _count_table(B, mem, (102, 109), "SD", "rebar", area=False)
        if g["kind"] == "steel":
            _steel_inputs(B, g, joints)  # 강재 표 끝(내화피복) 바로 다음이 기울기 및 침하
        elif i < len(groups) - 1:
            B.block(110, 110)
    # ---- 기울기 및 침하
    r = B.block(310, 315)
    tv = tests_view
    for i, key in enumerate(("tilt", "settlement")):
        g = (tv.get(key) or {}).get("grade") or ""
        t = DEFORM_TEXT.get(g, ("-", "-"))
        B.put("N", r + 2 + i, t[0])
        B.put("AN", r + 2 + i, t[1])
    return B.save(dest)


def _test_rows(B: Book, mem: dict, item: str, head: tuple, row: int, gap: int, empty: tuple, fmt) -> None:
    rows_by = {k: ((mem.get(k) or {}).get("items", {}).get(item) or {}).get("rows", []) for k in MEMBERS}
    n = max((len(v) for v in rows_by.values()), default=0)
    if not n:
        B.block(*empty)  # SafetyMan: 제목 + 머리(키 큰 줄) + 빈 줄
        return
    B.block(*head)
    for i in range(n):
        r = B.block(row, row)
        B.put("B", r, f"평가부재{i + 1}")
        B.clear(DS_COLS, r)
        for k, col in zip(MEMBERS, DS_COLS):
            rs = rows_by[k]
            if i < len(rs):
                B.put(col, r, fmt(rs[i]))
                c = B.dst[f"{col}{r}"]
                c.alignment = copy.copy(c.alignment) if c.alignment.wrap_text else Alignment(horizontal="center", vertical="center", wrap_text=True)
    B.block(gap, gap)


def _count_table(B: Book, mem: dict, rows: tuple, item: str, sub: str | None, area: bool) -> None:
    r0 = B.block(*rows)
    for i in range(5):
        B.clear(CNT_COLS, r0 + 2 + i)
    if area:
        B.clear(CNT_COLS, r0 + 7)
        B.clear(CNT_COLS, r0 + 8)
    for k, col in zip(MEMBERS, CNT_COLS):
        it = ((mem.get(k) or {}).get("items") or {}).get(item)
        if it and sub:
            it = (it.get("sub") or {}).get(sub)
        if not it:
            continue
        for i, g in enumerate(rating.GRADES):
            v = (it.get("counts") or {}).get(g) or 0
            if v:
                B.put(col, r0 + 2 + i, int(v))
        if area:
            tot = it.get("total") or sum((it.get("counts") or {}).values())
            if tot:
                B.put(col, r0 + 7, int(tot))


def _steel_slot(kind: str, row: dict) -> str:
    """강재 줄 → SafetyMan 강재 칸(기둥·큰보·작은보) — 재료시험 줄의 sub가 먼저, 없으면 부재 이름."""
    if row.get("sub") in STEEL_DS_COLS:
        return row["sub"]
    if kind in STEEL_DS_COLS:
        return kind
    return "기둥" if "기둥" in kind else "큰보"


def _steel_inputs(B: Book, g: dict, joints: dict | None = None) -> None:
    """철골조 묶음 — 강재규격(측정값/설계값 mm²) 평가부재 줄 + 용접·볼트·부식·내화피복 표. 용접·볼트·부식 개수는 상태평가 화면 입력(joints,
    그 층 {기둥|큰보|작은보: {weld_n, weld_g, bolt_bad, bolt_n, corr_n, corr_g}}) — 없으면 빈칸(10/15)."""
    by: dict[str, list] = {}
    for k, r in g["steel"]:
        by.setdefault(_steel_slot(k, r), []).append(r)
    n = max((len(v) for v in by.values()), default=0)
    B.block(281, 282)
    for i in range(n):
        r = B.block(283, 283)
        B.put("B", r, f"평가부재{i + 1}")
        B.clear(tuple(STEEL_DS_COLS.values()), r)
        for sub, col in STEEL_DS_COLS.items():
            rs = by.get(sub, [])
            if i < len(rs):
                x = rs[i]
                if x.get("area_measured") and x.get("area_design"):
                    B.put(col, r, f"{x['area_measured'] * 100:.0f} / {x['area_design'] * 100:.0f}")
                else:
                    B.put(col, r, f"{x.get('measured') or '-'} / {x.get('design') or '-'}")
    B.block(285, 285)
    jt = (joints or {}).get(str(g["agg"]["floor"])) or {}
    for a, b, rows in ((286, 293, range(288, 293)), (294, 297, range(296, 297)), (298, 305, range(300, 305))):
        r0 = B.block(a, b)
        for rr in rows:
            B.clear(tuple(STEEL_CNT_COLS.values()), r0 + rr - a)
        for m, col in STEEL_CNT_COLS.items():
            x = jt.get(m) or {}
            if a == 286 and x.get("weld_n"):  # 용접접합상태 — 상태 줄(a 매우양호 ~ e 매우불량)에 개수
                B.put(col, r0 + 2 + "abcde".index(x.get("weld_g") or "a"), int(x["weld_n"]))
            elif a == 294 and (x.get("bolt_n") or x.get("bolt_bad")):  # 볼트접합부 — 불량/전체(SafetyMan "0/20")
                B.put(col, r0 + 2, f"{int(x.get('bolt_bad') or 0)}/{int(x.get('bolt_n') or 0)}")
            elif a == 298 and x.get("corr_n"):  # 강재부식도 — 상태 줄(a 거의 없음 ~ e 15%이상)에 개수
                B.put(col, r0 + 2 + "abcde".index(x.get("corr_g") or "a"), int(x["corr_n"]))
    B.block(306, 309)
