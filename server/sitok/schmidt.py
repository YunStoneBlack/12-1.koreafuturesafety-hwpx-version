"""슈미트해머(반발경도) 시험성과표 엑셀 읽기(2026-10-13, 정밀안전점검 재료시험 — 형 메모 "슈미트 장비 결과 엑셀표를 첨부하면 자동 입력").

견본 `강도측정결과.xlsx`(바탕화면이 아니라 `Documents\\카카오톡 받은 파일\\정밀점검 레퍼런스\\`): Sheet1 "시험성과표" — 위치마다 4줄 × 5칸 = 측정치 20개,
H 평균(R1) · I ±20% · J 이상치 뺀 평균(R2) · K 타격각도(α) · L 각도보정(ΔR) · M 습윤보정(ΔR') · N 기준강도(R0) · O/P 압축강도 재료학회/건축학회 ·
Q 타설 년도(재령) · R 재령계수(n) · S/T 추정압축강도 재료학회/건축학회 · U Schmidt Hammer 추정강도(MPa, 둘의 평균) · V 비고. 값은 수식이 아니라 숫자로 들어 있음.
돌려주는 값은 위치마다 dict — 보고서 "콘크리트 강도조사" 표(위치·측정강도)와 rating.grade_strength(측정÷설계)에 씀.
"""
from __future__ import annotations

import io
from dataclasses import asdict, dataclass

import openpyxl

COLS = {"r1": 8, "pm20": 9, "r2": 10, "angle": 11, "angle_corr": 12, "wet_corr": 13, "r0": 14, "fc_mat": 15, "fc_arch": 16,
        "age_years": 17, "age_coef": 18, "est_mat": 19, "est_arch": 20, "est": 21, "note": 22}


@dataclass
class Spot:
    position: str  # HT-01, KS001 …
    readings: list[float]  # 측정치(보통 20개)
    r1: float | None = None
    pm20: float | None = None  # ±20%
    r2: float | None = None
    angle: float | None = None
    angle_corr: float | None = None
    wet_corr: float | None = None
    r0: float | None = None
    fc_mat: float | None = None
    fc_arch: float | None = None
    age_years: float | None = None
    age_coef: float | None = None
    est_mat: float | None = None
    est_arch: float | None = None
    est: float | None = None  # Schmidt Hammer 추정강도(MPa) — 보고서 측정강도
    note: str = ""


def _num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def read(data: bytes | str) -> list[Spot]:
    """엑셀(바이트 또는 경로) → 위치별 Spot. 시험성과표 시트가 없으면 첫 시트. 위치 칸(B열)이 있는 줄부터 4줄이 한 위치."""
    wb = openpyxl.load_workbook(io.BytesIO(data) if isinstance(data, bytes) else data, data_only=True)
    ws = next((s for s in wb.worksheets if any("시험성과표" in str(c.value or "") for row in s.iter_rows(max_row=3) for c in row)), wb.worksheets[0])
    header_row = next((r for r in range(1, 12) if any("측정치" in str(ws.cell(r, c).value or "") for c in range(1, 8))), 4)
    spots: list[Spot] = []
    for row in range(header_row + 1, ws.max_row + 1):
        pos = ws.cell(row, 2).value
        if pos in (None, "") or _num(ws.cell(row, 3).value) is None:
            continue
        readings = [v for rr in range(row, row + 4) for cc in range(3, 8) if (v := _num(ws.cell(rr, cc).value)) is not None]
        vals = {k: ws.cell(row, c).value for k, c in COLS.items()}
        spots.append(Spot(position=str(pos).strip(), readings=readings, note=str(vals.pop("note") or "").strip(),
                          **{k: _num(v) for k, v in vals.items()}))
    if not spots:
        raise ValueError("시험성과표에서 측정 위치를 찾지 못했습니다(B열 위치·C~G열 측정치).")
    return spots


def summary(spots: list[Spot]) -> dict:
    """보고서 결과분석에 쓰는 요약 — 추정강도 최소·최대·평균(MPa)."""
    ests = [s.est for s in spots if s.est is not None]
    return {"count": len(spots), "min": min(ests) if ests else None, "max": max(ests) if ests else None,
            "mean": round(sum(ests) / len(ests), 1) if ests else None, "spots": [asdict(s) for s in spots]}
