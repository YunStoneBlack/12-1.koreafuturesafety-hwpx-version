"""시특법 대가내역서(2026-10-15 민재형 — FMS e-보고서 "계약서 및 대가내역서" 폴더, 양식 = 형 엑셀 data/templates/시특법_대가내역서_양식.xlsx).

형: "대가내역서는 그 보고서에 올리는 총계약금액을 설명해야 함 — 묶음 계약이면 계약 총액과 이번 용역 범위의 대가가 보여야 반려가 안 남"
(예시 = 숭의 두 동 × 2026상 정밀 170만 + 정기 55만 5반기 = 890만). 보통은 한 시설물 상·하반기 두 줄.
- 줄 = 계약 회차 줄(civil.normalize_rounds). 동이 여럿이면 회차마다 동별 금액(parts) → 동마다 묶음(시설명 칸 세로 병합) + 반기 줄.
  동별 금액이 없으면 한 동(시설물명)에 회차 줄 그대로. 끝 없는 계약 = "매 반기" 한 줄(단가), 관급(회차 없음) = 계약기간 한 줄(계약금액).
- 양식 모양 그대로: 표 머리(10행) 위·합계·비고·회사 칸은 형 양식의 글·서식을 옮기고, 가운데 줄만 개수에 맞게 다시 그림(병합·테두리 복사).
- 엑셀 → PDF는 LibreOffice(server/contract_docs/to_pdf.py, 착수계와 같음).
"""
from __future__ import annotations

import copy
import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

from server.contract_docs.build import TEMPLATE_DIR
from server.sitok import civil

TEMPLATE = TEMPLATE_DIR / "시특법_대가내역서_양식.xlsx"
HEAD, FIRST, SUM_ROW, LAST_ROW = 10, 11, 23, 38  # 양식: 표 머리 10행, 줄 11~22, 합계 23행, 인쇄 끝 38행
COLS = 13  # A~M(인쇄 영역)


def blocks(c) -> list[tuple[str, list[tuple[str, str, int]]]]:
    """계약 → [(시설명, [(과업기간, 분야, 금액)])]."""
    name = c.facility_name or c.client or ""
    rounds = civil.normalize_rounds(c.rounds)
    if getattr(c, "open_ended", False):
        u = c.unit or {}
        return [(name, [("매 반기", u.get("kind") or "정기", int(u.get("amount") or 0))])]
    label = lambda r: f"{r['year']}년{r['half']}"  # noqa: E731 — 양식 글 모양 "2026년상반기"
    names = civil.buildings(rounds)
    if names:
        out = []
        for n in names:
            rows = [(label(r), r["kind"], civil.part_amount(r, n)) for r in rounds]
            out.append((n, [(a, k, v) for a, k, v in rows if v is not None]))
        return out
    if rounds:
        return [(name, [(label(r), r["kind"], r["amount"]) for r in rounds])]
    s, e = getattr(c, "start_date", None), getattr(c, "end_date", None)
    period = f"{s:%Y.%m.%d}~{e:%Y.%m.%d}" if s and e else ""
    return [(name, [(period, "정기", int(c.amount or 0))])]


def _style(cell) -> dict:
    return {"font": copy.copy(cell.font), "border": copy.copy(cell.border), "fill": copy.copy(cell.fill),
            "number_format": cell.number_format, "alignment": copy.copy(cell.alignment), "protection": copy.copy(cell.protection)}


def _put_style(cell, st: dict) -> None:
    for k, v in st.items():
        setattr(cell, k, v)


def build(c, dest: Path, day: datetime.date | None = None) -> Path:
    """대가내역서 엑셀을 dest에 만든다(c = SitokDocContract 같은 모양 — facility_name·client·title·doc_no·rounds·unit·open_ended·amount·contract_date)."""
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    # 양식에서 옮길 것: 줄 서식(묶음 첫 줄·가운데·끝), 합계~끝 줄(글·서식·높이·병합)
    # 묶음 첫 줄 서식: 맨 첫 묶음 = 11행(표 머리 아래 굵은 선), 다음 묶음부터 = 17행(가는 선)
    rows_of = (("first0", 11), ("first", 17), ("mid", 12), ("last", 16))
    row_st = {k: [_style(ws.cell(r, col)) for col in range(1, COLS + 1)] for k, r in rows_of}
    row_h = {k: ws.row_dimensions[r].height for k, r in rows_of}
    tail = [(r, [(ws.cell(r, col).value, _style(ws.cell(r, col))) for col in range(1, COLS + 1)], ws.row_dimensions[r].height) for r in range(SUM_ROW, LAST_ROW + 1)]
    tail_merges = [m for m in ws.merged_cells.ranges if m.min_row >= SUM_ROW and m.max_col <= COLS]
    tail_merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in tail_merges]
    for m in list(ws.merged_cells.ranges):  # 줄 칸·오른쪽 계산용 칸(N열부터 — 양식 작업용)의 병합은 다 풂
        if m.min_row >= FIRST or m.min_col > COLS:
            ws.unmerge_cells(str(m))
    ws.delete_rows(FIRST, ws.max_row - FIRST + 1)
    for row in ws.iter_rows(min_row=1, max_row=FIRST - 1, min_col=COLS + 1, max_col=ws.max_column):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None

    r = FIRST
    data = blocks(c)
    for b, (name, rows) in enumerate(data):
        start = r
        for i, (period, kind, amount) in enumerate(rows):
            pos = ("first0" if b == 0 else "first") if i == 0 else ("last" if i == len(rows) - 1 else "mid")
            for col in range(1, COLS + 1):
                _put_style(ws.cell(r, col), row_st[pos][col - 1])
            ws.row_dimensions[r].height = row_h[pos] or 24.75
            ws.cell(r, 4, period)
            ws.cell(r, 5, kind)
            ws.cell(r, 7, amount)
            ws.cell(r, 10, f"=G{r}")
            for a, b in ((5, 6), (7, 9), (10, 11), (12, 13)):
                ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            r += 1
        ws.cell(start, 1, name)
        ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=3)
    sum_row = r
    shift = sum_row - SUM_ROW
    for old, cells, h in tail:
        nr = old + shift
        for col, (v, st) in enumerate(cells, start=1):
            cell = ws.cell(nr, col)
            _put_style(cell, st)
            cell.value = v if not (isinstance(v, str) and v.startswith("=")) else None  # 양식 수식(다른 칸 참조)은 아래에서 값으로
        ws.row_dimensions[nr].height = h
    for r1, c1, r2, c2 in tail_merges:
        ws.merge_cells(start_row=r1 + shift, start_column=c1, end_row=r2 + shift, end_column=c2)
    ws.cell(sum_row, 10, f"=SUM(J{FIRST}:K{sum_row - 1})")
    total = sum(a for _, rows in data for _, _, a in rows)
    # 위쪽 칸
    d = getattr(c, "contract_date", None) or day or datetime.date.today()
    ws["B2"] = c.doc_no or f"KFSC_{d:%Y}_{d:%m%d}"
    ws["B6"] = c.title or civil.default_title(c)
    ws["B8"] = " 금"  # 양식은 오른쪽 계산 칸(V4·V5)을 참조 — 그 칸을 지워서 글로
    ws["C8"] = total
    ws["D8"] = " 원정"
    ws["F8"] = total
    for img in ws._images:  # 도장 그림(32~35행)도 줄 수만큼 같이 옮김
        img.anchor._from.row += shift
        if getattr(img.anchor, "to", None) is not None:
            img.anchor.to.row += shift
    ws.cell(29 + shift, 6, d)
    who = c.client or c.facility_name or ""
    ws.cell(37 + shift, 1, f"{who} 시설물 안전점검 담당자님" if who else "시설물 안전점검 담당자님")
    ws.print_area = f"A1:M{LAST_ROW + shift}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def build_pdf(c, xlsx: Path, pdf: Path) -> Path:
    from server.contract_docs.to_pdf import office_to_pdf

    build(c, xlsx)
    return office_to_pdf(xlsx, pdf)
