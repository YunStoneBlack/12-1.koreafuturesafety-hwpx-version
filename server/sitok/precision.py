"""시특법 정밀안전점검 ④ — 정기에 없는 입력(2026-10-13 밤, 사용자 "추천대로"): 2장 자료수집(2.2 자료 목록·2.3.2 내진·2.3.4 사용·관리실태·전차 점검 보고서 검토),
3.2.6 하중상태 사진 4장, 3.2.7~8 공중이용·기타시설 등급. 항목·문구는 다산은혜교회 26년 상반기 보고서 그대로. 저장은 SitokReport.survey["precision"](마이그레이션 없음),
등급은 SitokReport.parts["grades"]. 화면 sitok-field.html [📋 자료수집]·[⚖ 하중상태] 탭 + 공중이용·기타시설 탭의 등급 칸(js/sitok-precision.js), 서버 sitok_precision.py.
"""
from __future__ import annotations

import io

# 2.2 자료수집 목록(본문 표 2.1 + 부록 사전검토보고서 4.2 유지관리자료 보유 현황 — 둘이 미묘하게 달라 합친 한 목록, 사용자 10/15 "한 번 고르면 둘 다") — (구분, 키, 자료)
DOC_ITEMS = [
    ("시공관련자료", "build", "시공관련 자료"), ("시공관련자료", "quality", "품질관리자료 — 재료증명서·품질시험기록·관리 및 선정시험 기록"),
    ("시공관련자료", "measure", "주요 구조부위 계측 관련자료"), ("시공관련자료", "accident", "사고기록"), ("시공관련자료", "operate", "운영기록"),
    ("설계도서", "spec", "공통 — 준공내역서·공사시방서·각종 계산서·토질 및 지반조사 보고서"),
    ("설계도서", "dwg_common", "공통도면 — 위치도(배치도)·평면도·단면도(종·횡)"), ("설계도서", "dwg_arch", "건축도면 — 배치도·평면도·입면도·단면도"),
    ("설계도서", "dwg_struct", "구조도면 — 평면·단면도·배근도·접합 상세도"), ("설계도서", "dwg_mech", "기계설비도면 — 승강기·냉난방·환기"),
    ("설계도서", "dwg_elec", "전기설비도면 — 조명·통신·방송·변전"), ("설계도서", "dwg_fire", "소방설비도면 — 방화구획·소화전·스프링클러"),
    ("설계도서", "dwg_water", "급배수설비도면 — 계통도·수조·정화조"),
    ("유지관련자료", "ledger", "시설물관리대장 — 기본현황·상세제원·유지관리 이력"), ("유지관련자료", "past", "기존 안전점검·정밀안전진단 자료"),
    ("유지관련자료", "repair", "보수·보강 및 용도변경 자료"),
]
DOC_GROUPS = ("시공관련자료", "설계도서", "유지관련자료")


def doc_groups(docs: list[dict]) -> list[dict]:
    """자료수집 목록 → 표 줄(구분마다) [{group, items: ["1) 자료", …], status: 보유/미보유/일부 보유/-, detail: ["자료 : 보유(메모)", …]}] — 본문 표 2.1·부록 4.2 공용."""
    out = []
    for g in DOC_GROUPS:
        rows = [d for d in docs if d["group"] == g]
        sts = [d["status"] for d in rows if d.get("status")]
        status = "-" if not sts else sts[0] if all(x == sts[0] for x in sts) else "일부 보유"
        detail = [f"{d['name'].split(' — ')[0]} : {d['status']}" + (f"({d['note']})" if d.get("note") else "") for d in rows if d.get("status")]
        out.append({"group": g, "items": [f"{i}) {d['name'].replace(' — ', ' : ')}" for i, d in enumerate(rows, 1)], "status": status, "detail": detail or ["-"]})
    return out


DOC_STATUS = ("보유", "미보유", "해당없음")
SEISMIC = ("적용", "미적용", "불명")
YNU = ("무", "유", "불명")
# 2.3.4 4) 부대 점검사항(유지관리자 면담) — 건축분야 25항목, 값 "없다(양호)" / "있다(불량)"
CHECKLIST = [
    "바닥 포장부위 침하 및 균열현상", "건물 전체의 부동침하 현상", "외부 옹벽(축대)의 균열현상", "하수관로 및 맨홀의 배수·청소 상태", "외벽의 전도위험 부위",
    "옥외 간판의 고정 상태", "외벽 타일의 탈락 부위", "외벽 모르터 또는 콘크리트의 탈락 부위", "외벽 창문 유리의 파손", "옥상 ROOF DRAIN의 상태",
    "옥상에 하중(물건)의 과적 여부", "옥상과 파라펫의 균열", "외벽 벽돌 줄눈의 경사방향 균열", "외벽의 백화 현상", "내부 창·문의 작동 상태",
    "건물 내부의 진동 여부", "천정재(텍스류)의 탈락 및 갈라짐 상태", "벽지 및 천정지가 찢어진 곳 유무", "실내의 하중(물건)의 과적 여부", "건물에서 뚝뚝하는 소리",
    "누수되는 곳의 유무", "녹물이 흘러나오는 곳의 유무", "코킹이 갑자기 떨어진 곳의 유무", "외부 석재에 균열 또는 변형 유무", "슬래브·보·기둥·벽체의 들뜬 현상",
]
LOAD_SLOTS = 4  # 3.2.6 사용하중 현황사진 4장(다산: 옥상·식당·믿음실·지하주차장)
# 조사 전경 사진(형 10/14 "조사 전경 모두 2개씩, 탄산화만 수직재 2·수평재 2") — (키, 보고서 캡션, 화면 제목, [칸 제목…]) — 공통편 2장 【사진 2.1】·시설물편 4장 【사진 4.1~4.5】
SURVEY_PICS = [
    ("visual", "【사진 2.1】", "외관조사 전경", ["육안관찰 외관조사", "육안관찰 외관조사"]),
    ("visual2", "【사진 3.1】", "시설물 외관조사 전경(정밀육안조사)", ["정밀육안조사", "정밀육안조사"]),
    ("section", "【사진 4.1】", "부재규격 조사 전경", ["부재 규격 조사(수직재)", "부재 규격 조사(수평재)"]),
    ("strength", "【사진 4.2】", "콘크리트 강도조사 전경", ["콘크리트 강도조사(수직재)", "콘크리트 강도조사(수평재)"]),
    ("carb", "【사진 4.3】", "탄산화시험 조사 전경", ["수직재 탄산화 시험(천공)", "수직재 탄산화 시험(피복두께 측정)", "수평재 탄산화 시험(천공)", "수평재 탄산화 시험(피복두께 측정)"]),
    ("tilt", "【사진 4.4】", "기울기(수평변위) 조사 전경", ["기울기(수평변위) 측정전경-1", "기울기(수평변위) 측정전경-2"]),
    ("settle", "【사진 4.5】", "부등침하(수직변위) 조사 전경", ["부등침하(수직변위) 측정전경-1", "부등침하(수직변위) 측정전경-2"]),
]
SURVEY_KEYS = {k: len(titles) for k, _, _, titles in SURVEY_PICS}
# 2.3.4 4) 부대 점검사항 자동(사용자 10/14 "결함 사진 분석해도 없는 거지?" → 결함 목록으로 해당 줄 "있다") — (체크리스트 번호, 판정 함수(결함) → 맞으면 True)
CHECK_RULES = {
    6: lambda d: "탈락" in (d.dtype or "") and any(w in (d.member or "") for w in ("타일", "외벽")),
    7: lambda d: any(w in (d.dtype or "") for w in ("박락", "박리", "재료분리")) and (d.part == "구조체" or "외벽" in (d.member or "")),
    11: lambda d: "옥상" in (d.floor or "") and "균열" in (d.dtype or ""),
    13: lambda d: any(w in (d.dtype or "") for w in ("백태", "백화")),
    16: lambda d: any(w in (d.member or "") for w in ("천장", "천정")) and any(w in (d.dtype or "") for w in ("탈락", "균열", "갈라")),
    20: lambda d: "누수" in (d.dtype or ""),
    21: lambda d: "부식" in (d.dtype or "") or "녹" in (d.dtype or ""),
    23: lambda d: "석재" in (d.member or "") and any(w in (d.dtype or "") for w in ("균열", "변형")),
    24: lambda d: d.part == "구조체" and any(w in (d.dtype or "") for w in ("들뜸", "층분리")),
}


def auto_checklist(defects: list) -> dict:
    """결함 목록 → {번호(str): {"types": [결함유형…], "where": "층 부재 …"}} — 점검자가 표에서 고친 것이 없을 때 "있다"로."""
    out: dict = {}
    for d in defects:
        for i, rule in CHECK_RULES.items():
            try:
                hit = rule(d)
            except Exception:  # noqa: BLE001
                hit = False
            if hit:
                e = out.setdefault(str(i), {"types": [], "where": []})
                if d.dtype and d.dtype not in e["types"]:
                    e["types"].append(d.dtype)
                w = f"{d.floor} {d.member}".strip()
                if w and w not in e["where"] and len(e["where"]) < 4:
                    e["where"].append(w)
    return {k: {"types": v["types"][:3], "where": ", ".join(v["where"])} for k, v in out.items()}
LOAD_DEFAULT_NOTE = "용도변경 이력은 없는 것으로 확인되었고, 현장에서도 추가하중에 대한 확인결과 준공도서와 동일한 것으로 조사되었다."
GRADES = ("a", "b", "c", "d", "e")


def empty() -> dict:
    return {"docs": {}, "seismic": {"status": "", "note": ""}, "manage": {"use_change": "무", "struct_change": "무", "surround_change": "무", "checklist": {}, "note": ""},
            "prev_review": {"text": ""}, "history": [], "load": {"photos": [], "note": ""}, "survey": {k: [""] * n for k, n in SURVEY_KEYS.items()}}


def norm(p: dict | None) -> dict:
    """저장 모양 맞추기(모르는 키·값은 버림)."""
    p = p or {}
    out = empty()
    keys = {k for _, k, _ in DOC_ITEMS}
    for k, v in (p.get("docs") or {}).items():
        if k in keys and isinstance(v, dict):
            out["docs"][k] = {"status": v.get("status") if v.get("status") in DOC_STATUS else "", "note": str(v.get("note") or "").strip()}
    s = p.get("seismic") or {}
    out["seismic"] = {"status": s.get("status") if s.get("status") in SEISMIC else "", "note": str(s.get("note") or "").strip()}
    m = p.get("manage") or {}
    out["manage"] = {k: (m.get(k) if m.get(k) in YNU else "무") for k in ("use_change", "struct_change", "surround_change")}
    out["manage"]["checklist"] = {str(i): bool(v) for i, v in (m.get("checklist") or {}).items() if str(i).isdigit() and int(i) < len(CHECKLIST)}
    out["manage"]["note"] = str(m.get("note") or "").strip()
    out["prev_review"] = {"text": str((p.get("prev_review") or {}).get("text") or "").strip()}
    out["history"] = [h for h in (p.get("history") or []) if isinstance(h, dict)]
    lo = p.get("load") or {}
    photos = []
    for ph in (lo.get("photos") or [])[:LOAD_SLOTS]:
        if isinstance(ph, dict):
            photos.append({"file": str(ph.get("file") or ""), "title": str(ph.get("title") or "").strip()})
    out["load"] = {"photos": photos, "note": str(lo.get("note") or "").strip()}
    sv = p.get("survey") or {}
    for k, n in SURVEY_KEYS.items():  # 조사 전경 사진 — 칸 수 고정, 값은 회차 폴더 상대경로
        got = [str(x or "") for x in (sv.get(k) or [])][:n]
        out["survey"][k] = got + [""] * (n - len(got))
    return out


def read_fms(data: bytes) -> list[dict]:
    """FMS 점검진단실적 엑셀(전회차 보고서.xlsx) → [{no, kind, start, end, org, chief, cost, grade, result, repair}] — 머리글 "점검진단구분" 줄 아래."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.worksheets[0]
    head = None
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [str(ws.cell(r, c).value or "").replace(" ", "") for c in range(1, ws.max_column + 1)]
        if any("점검진단구분" in v for v in vals):
            head = (r, vals)
            break
    if head is None:
        raise ValueError("점검진단실적 양식이 아닙니다(머리글 '점검진단구분'이 없음).")
    r0, vals = head
    col = {}
    for i, v in enumerate(vals, 1):
        for key, words in (("kind", ("점검진단구분",)), ("start", ("시작일",)), ("end", ("종료일",)), ("org", ("점검진단기관",)), ("chief", ("책임기술자",)),
                           ("cost", ("비용",)), ("grade", ("안전등급",)), ("result", ("주요점검진단결과",)), ("repair", ("주요보수보강안",))):
            if key not in col and any(w in v for w in words):
                col[key] = i
    rows = []
    for r in range(r0 + 1, ws.max_row + 1):
        get = lambda k: str(ws.cell(r, col[k]).value or "").strip() if k in col else ""  # noqa: E731
        if not get("kind") and not get("start"):
            continue
        rows.append({"no": len(rows) + 1, "kind": get("kind"), "start": get("start")[:10], "end": get("end")[:10], "org": get("org"), "chief": get("chief"),
                     "cost": get("cost"), "grade": get("grade").replace("등급", ""), "result": get("result"), "repair": get("repair")})
    if not rows:
        raise ValueError("점검진단실적에 줄이 없습니다.")
    return rows


def _selfcheck() -> None:
    assert len(CHECKLIST) == 25 and len(DOC_ITEMS) == 16
    n = norm({"docs": {"ledger": {"status": "보유", "note": "컴퓨터 파일"}, "zzz": {"status": "보유"}}, "manage": {"use_change": "유", "checklist": {"3": True, "99": True}},
              "load": {"photos": [{"file": "a.jpg", "title": "옥상"}], "note": ""}})
    assert list(n["docs"]) == ["ledger"] and n["manage"]["use_change"] == "유" and n["manage"]["checklist"] == {"3": True} and n["load"]["photos"][0]["title"] == "옥상"
    rows = read_fms(open(r"C:\Users\윤석현1\Desktop\시특법 견본\정밀\전회차 보고서.xlsx", "rb").read())
    assert len(rows) == 3 and rows[0]["grade"] == "C" and rows[2]["org"] == "(주)한국미래안전건설" and rows[0]["start"] == "2022-10-04", rows[0]
    print("정밀 입력 검증 통과")


if __name__ == "__main__":
    import sys

    sys.stdout.reconfigure(encoding="utf-8")
    _selfcheck()
