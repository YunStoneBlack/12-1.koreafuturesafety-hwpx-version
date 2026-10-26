"""내진성능평가 철근탐사(2026-10-21 — 2단계) — 회차 survey["rebar"] = {"rows": [...]}.

한 줄 = 시험 하나(견본 표: 본문 2.2.2 철근배근 조사결과 = 부록 "철근배근 상태조사", 같은 표):
  시험번호 · 측정위치 "지상1층( RC-1 )" · 부재명 "기둥 중앙부 단변 (C1)" · 두 줄 [주근 5EA→5EA 피복 40] [대근 @250→@243 피복 44].
  기둥 = 주근 개수(한 변 — 단변·장변)·대근 간격(단부·중앙부), 보 = 하부근 개수(두 줄이면 "nEA(복)")·늑근 간격, 벽체 = 수직근·수평근 간격.
  못 잰 줄 = "실측불가"(피복 "-").
설계값은 구조 일람표 부재 목록(server/sitok/rebar_schedule — 시설물 ledger.rebar_schedule)에서 층·부호·위치·변으로 자동(사람이 고치면 그대로).
조사값·피복은 탐사기 캡처를 올리면 server/sitok/rebar_capture가 채움(간격 = 철근 사이 평균, 피복 = 평균 반올림).
필요 수량(매뉴얼 <표 3-1>): 조사단위(층)마다 부재 종별(보·기둥·보-기둥 접합부·벽체) 3개소 이상 — 표본 층 × 종별 칸을 미리 만들어 둠(apply_plan, 10/22).
보-기둥 접합부는 종별 항목만 따로, 조사는 보 단부로(형 10/22 — 견본 "접합부 철근배근상태 조사는 보단부 조사로 갈음"): 일람표는 보 단부 칸, 줄은 하부근·늑근.
결과(조사값·실측불가·캡처)가 없는 미리 만든 칸은 수량·보고서에서 뺌(report_rows — 번호 다시).
"""
from __future__ import annotations

import io
import re
import uuid

from server.sitok import rebar_schedule as RS

JOINT = "보-기둥 접합부"
KINDS = ("보", "기둥", JOINT, "벽체")  # 철근배근 상태조사 종별(매뉴얼 <표 3-1> 비고) — 칸 차례도 이대로
LINES = {"기둥": ("주근", "대근"), "보": ("하부근", "늑근"), JOINT: ("하부근", "늑근"), "벽체": ("수직근", "수평근")}
COUNT_LINE = {"기둥": True, "보": True, JOINT: True, "벽체": False}  # 첫 줄이 개수(EA)인지(벽체는 둘 다 간격)
NEED_PER = 3


def sched_kind(kind: str) -> str:
    """일람표에서 찾을 부재 — 접합부는 보(단부 칸)."""
    return "보" if kind == JOINT else kind


def _int(v):
    try:
        n = int(round(float(v)))
        return n if n >= 0 else None
    except (TypeError, ValueError):
        return None


def norm_row(raw: dict, n: int) -> dict | None:
    kind = raw.get("kind") if raw.get("kind") in LINES else None
    if kind is None:
        return None
    floor = raw.get("floor")
    floor = int(floor) if isinstance(floor, (int, float)) or (isinstance(floor, str) and re.fullmatch(r"-?\d+", floor or "")) else None

    def pair(k, cast=str):
        v = list(raw.get(k) or [])[:2]
        v += [None] * (2 - len(v))
        return [cast(x) if x not in (None, "") else (None if cast is not str else "") for x in v]

    return {"id": raw.get("id") or uuid.uuid4().hex[:8], "no": n, "floor": floor, "kind": kind, "item_id": str(raw.get("item_id") or ""),
            "mark": str(raw.get("mark") or "").strip()[:30],
            "position": "" if kind == "벽체" else "단부" if kind == JOINT else (raw.get("position") if raw.get("position") in ("단부", "중앙부") else "중앙부"),  # 벽체는 위치 없음, 접합부는 보 단부
            "side": (raw.get("side") if raw.get("side") in ("단변", "장변") else "단변") if kind == "기둥" else "",
            "design": pair("design"), "design_manual": bool(raw.get("design_manual")),
            "measured": pair("measured"), "cover": pair("cover", _int), "unable": [bool(x) for x in (list(raw.get("unable") or []) + [False, False])[:2]],
            "captures": [str(x or "") for x in (list(raw.get("captures") or []) + ["", ""])[:2]],
            "warn": [str(x) for x in (raw.get("warn") or [])][:4], "note": str(raw.get("note") or "")[:60]}


def normalize(data: dict | None) -> dict:
    rows = []
    for raw in (data or {}).get("rows") or []:
        r = norm_row(raw, len(rows) + 1)
        if r:
            rows.append(r)
    return {"rows": rows, "seeded": bool((data or {}).get("seeded"))}  # seeded = 표본 칸을 한 번 깔았음


def items_by_id(schedule: dict | None) -> dict:
    return {x["id"]: x for x in RS.view(schedule).get("items") or []}


def choices(schedule: dict | None, floor: int | None, kind: str) -> list[dict]:
    """그 층·부재에 맞는 일람표 칸 [{id, label, position}] — 층 표기를 못 알아들은 칸·층을 안 고른 경우는 다 보임."""
    out = []
    for it in RS.view(schedule).get("items") or []:
        if it["kind"] != sched_kind(kind) or (kind == JOINT and "중앙" in (it.get("position") or "")):  # 접합부 = 보 단부 칸
            continue
        rng = RS.floor_range(it["floors"])
        if floor is not None and rng and not (rng[0] <= floor <= rng[1]):
            continue
        label = ", ".join(it["marks"]) + (f" {it['position']}" if it.get("position") else "") + f" ({it['floors']})"
        out.append({"id": it["id"], "label": label, "mark": ", ".join(it["marks"]), "position": it.get("position") or "", "confirmed": it["confirmed"],
                    "check_ok": (it.get("check") or {}).get("ok")})
    return out


def fill_design(row: dict, schedule: dict | None) -> dict:
    """설계값 자동(사람이 고친 줄은 그대로) — 기둥은 단변/장변·단부/중앙부, 보는 그 칸(위치), 벽은 수직·수평."""
    if row.get("design_manual"):
        return row
    it = items_by_id(schedule).get(row.get("item_id") or "")
    if it is None:
        return row
    pos = row.get("position") or ""
    if it["kind"] == "보" and it.get("position") and row["kind"] != JOINT:  # 접합부는 늘 단부
        pos = "단부" if "단부" in it["position"] else "중앙부" if "중앙" in it["position"] else pos
        row["position"] = pos
    rows = RS.design(it, pos, row.get("side") or "")["rows"]
    row["design"] = [rows[0]["value"], rows[1]["value"]]
    row["mark"] = row.get("mark") or ", ".join(it["marks"])
    return row


def apply_capture(row: dict, line: int, got: dict) -> dict:
    """캡처 읽은 값 → 그 줄 조사값·피복. 개수 줄(기둥 주근·보 하부근)은 개수, 간격 줄은 간격. 보 두 줄 설계면 "(복)"도 붙임."""
    is_count = line == 0 and COUNT_LINE[row["kind"]]
    if is_count and got.get("count"):
        dup = "(복)" if "(복)" in (row["design"][0] or "") else ""
        row["measured"][0] = f"{got['count']}EA{dup}"
    elif not is_count and got.get("spacing"):
        row["measured"][line] = f"@{got['spacing']}"
    if got.get("cover") is not None:
        row["cover"][line] = got["cover"]
    row["unable"][line] = False
    row["warn"] = list(dict.fromkeys([*row.get("warn", []), *[f"{LINES[row['kind']][line]}: {w}" for w in got.get("warn") or []]]))[:4]
    return row


def mismatch(row: dict) -> list[int]:
    """설계와 조사가 다른 줄(본문 "B1부재가 구조도면과 실측 수치와 상이" 글 재료) — 개수는 다르면, 간격은 설계의 ±20% 밖이면."""
    out = []
    for i in (0, 1):
        d, m = row["design"][i] or "", row["measured"][i] or ""
        if not d or not m or row["unable"][i]:
            continue
        dn, mn = re.search(r"\d+", d), re.search(r"\d+", m)
        if not dn or not mn:
            continue
        a, b = int(dn.group()), int(mn.group())
        if "@" in d:
            if a and abs(b - a) / a > 0.2:
                out.append(i)
        elif a != b:
            out.append(i)
    return out


def has_result(row: dict) -> bool:
    """조사 결과가 있는 줄(보고서·수량에 셈) — 조사값·실측불가·캡처 중 하나."""
    return any(row["measured"]) or any(row["unable"]) or any(row["captures"])


def touched(row: dict) -> bool:
    """사람이 손댄 줄(칸 수를 줄여도 안 지움) — 일람표 칸·부호·결과·비고·설계 직접."""
    return bool(row.get("item_id") or row.get("mark") or row.get("note") or row.get("design_manual") or has_result(row))


def report_rows(rows: list[dict]) -> list[dict]:
    """보고서·엑셀에 넣을 줄 — 결과 있는 줄만, 시험번호 1부터 다시."""
    return [{**r, "no": i} for i, r in enumerate((x for x in rows if has_result(x)), 1)]


def kind_short(row: dict) -> str:
    """부재 칸에 쓸 이름 — 접합부는 보(보 단부로 조사)."""
    return sched_kind(row["kind"])


def default_plan(floors: list[int]) -> dict:
    """표본 층 × 4종 × 3개소."""
    return {f"{n}|{k}": NEED_PER for n in floors for k in KINDS}


def apply_plan(rows: list[dict], plan: dict) -> list[dict]:
    """칸 수 {"층|종별": n}대로 줄 맞추기 — 모자라면 빈 줄 더하고, 남으면 손 안 댄 줄만 뒤에서부터 지움.
    층(아래 → 위)·종별 차례로 세우고 번호 다시. plan에 없는 층·종별 줄은 그대로."""
    order = {k: i for i, k in enumerate(KINDS)}
    groups: dict[tuple, list] = {}
    for r in rows:
        groups.setdefault((r["floor"], r["kind"]), []).append(r)
    want = {}
    for key, n in (plan or {}).items():
        m = re.fullmatch(r"(-?\d+)\|(.+)", str(key))
        if m and m.group(2) in LINES and int(m.group(1)):
            want[(int(m.group(1)), m.group(2))] = max(0, min(_int(n) or 0, 40))
    out = []
    for key in sorted(set(groups) | set(want), key=lambda k: (k[0] is None, k[0] or 0, order.get(k[1], 9))):
        g = list(groups.get(key, []))
        if key in want:
            i = len(g) - 1
            while len(g) > want[key] and i >= 0:
                if not touched(g[i]):
                    g.pop(i)
                i -= 1
            while len(g) < want[key]:
                g.append(norm_row({"kind": key[1], "floor": key[0]}, 0))
        out.extend(g)
    for i, r in enumerate(out, 1):
        r["no"] = i
    return out


def resync_floors(rows: list[dict], old: list[int], new: list[int]) -> list[dict] | None:
    """조사단위를 바꿔 표본 층이 달라지면(10/26) — 새 층은 4종 × 3칸(그 층 줄이 이미 있으면 그대로), 빠진 층은 손 안 댄 줄만 지움(쓴 줄은 둠). 그대로면 None."""
    add, drop = [n for n in new if n not in old], [n for n in old if n not in new]
    if not add and not drop:
        return None
    have = {r["floor"] for r in rows}
    plan = {f"{n}|{k}": NEED_PER for n in add if n not in have for k in KINDS}
    plan.update({f"{n}|{k}": 0 for n in drop for k in KINDS})
    return apply_plan(rows, plan)


def counts(rows: list[dict], floors: list[int]) -> dict:
    """수량 표 — 층마다 종별 {결과 있는 줄(done), 칸(slots)}, 합계(결과 있는 줄·칸)."""
    fl = sorted(set(floors) | {r["floor"] for r in rows if r["floor"] is not None})
    done = {f: {k: 0 for k in KINDS} for f in fl}
    slots = {f: {k: 0 for k in KINDS} for f in fl}
    for r in rows:
        if r["floor"] in slots:
            slots[r["floor"]][r["kind"]] += 1
            done[r["floor"]][r["kind"]] += 1 if has_result(r) else 0
    return {"per": NEED_PER, "kinds": list(KINDS), "table": [{"floor": f, "label": RS.floor_label(f), "done": done[f], "slots": slots[f]} for f in fl],
            "total": sum(1 for r in rows if has_result(r)), "slots_total": len(rows)}


def place(row: dict) -> str:
    return f"{RS.floor_label(row['floor']) if row['floor'] is not None else '층?'}( RC-{row['no']} )"


def member_text(row: dict) -> tuple[str, str]:
    """부재명 두 줄 — "기둥 중앙부 단변" / "(C1)". 접합부는 "보 단부(접합부)"."""
    head = "보 단부(접합부)" if row["kind"] == JOINT else " ".join(x for x in (row["kind"], row["position"], row["side"]) if x)
    return head, f"({row['mark']})" if row["mark"] else ""


def table_rows(rows: list[dict]) -> list[list]:
    """표 줄(두 줄씩) — [시험번호, 측정위치, 부재명, 구분, 설계도면, 조사결과, 피복]."""
    out = []
    for r in rows:
        head, mark = member_text(r)
        for i in (0, 1):
            meas = "실측불가" if r["unable"][i] else (r["measured"][i] or "")
            cov = "-" if r["unable"][i] else (r["cover"][i] if r["cover"][i] is not None else "")
            out.append([r["no"] if i == 0 else "", place(r) if i == 0 else "", head if i == 0 else mark, LINES[r["kind"]][i], r["design"][i] or "", meas, cov])
    return out


def xlsx(rows: list[dict], title: str) -> bytes:
    """철근배근 상태조사 표(엑셀) — 견본 모양: 시험번호·측정위치는 두 줄 합침, 단위 mm."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "철근배근 상태조사"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    mid = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.append([title])
    ws.append(["(단위 : mm)"])
    ws.append(["시험\n번호", "측정위치", "부재명", "설계도면", "", "조사결과", "피복\n(mm)"])
    ws.merge_cells("D3:E3")
    for c in ws[3]:
        c.font, c.alignment, c.border, c.fill = Font(bold=True), mid, box, PatternFill("solid", fgColor="EDEFF3")
    for k, line in enumerate(table_rows(rows)):
        ws.append(line)
        rr = ws.max_row
        for c in ws[rr]:
            c.alignment, c.border = mid, box
        if k % 2 == 1:
            for col in ("A", "B"):
                ws.merge_cells(f"{col}{rr - 1}:{col}{rr}")
    for col, w in zip("ABCDEFG", (7, 18, 18, 8, 12, 12, 9)):
        ws.column_dimensions[col].width = w
    ws["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _selfcheck() -> None:
    sch = {"items": [
        {"id": "c1", "kind": "기둥", "marks": ["C1"], "floors": "1~4", "size": "400×500", "main": "10-HD19", "short_side": 3, "long_side": 4,
         "hoop": "HD10@250", "hoop_end": "HD10@125", "confirmed": True},
        {"id": "g1", "kind": "보", "marks": ["G1"], "floors": "2~5", "position": "중앙부", "bottom": "6-HD19", "bottom_row": 3, "bottom_layers": 2, "stirrup": "HD10@300", "confirmed": True},
        {"id": "w1", "kind": "벽체", "marks": ["W1"], "floors": "PIT층/지하층", "thickness": 200, "vertical": "HD13@200", "horizontal": "HD10@250"}]}
    assert [c["id"] for c in choices(sch, 1, "기둥")] == ["c1"] and choices(sch, 6, "기둥") == [] and [c["id"] for c in choices(sch, -1, "벽체")] == ["w1"]
    r = fill_design(norm_row({"kind": "기둥", "floor": 1, "item_id": "c1", "position": "단부", "side": "단변"}, 1), sch)
    assert r["design"] == ["3EA", "@125"] and r["mark"] == "C1"
    r = apply_capture(r, 0, {"count": 3, "cover": 51, "warn": []})
    r = apply_capture(r, 1, {"spacing": 130, "cover": 62, "warn": []})
    assert r["measured"] == ["3EA", "@130"] and r["cover"] == [51, 62] and mismatch(r) == []
    b = fill_design(norm_row({"kind": "보", "floor": 2, "item_id": "g1"}, 2), sch)
    assert b["design"] == ["3EA(복)", "@300"] and b["position"] == "중앙부"
    b = apply_capture(b, 0, {"count": 3, "cover": 63, "warn": []})
    assert b["measured"][0] == "3EA(복)"
    b["measured"][1] = "@400"
    assert mismatch(b) == [1]
    t = table_rows([r, b])
    assert t[0][:3] == [1, "지상1층( RC-1 )", "기둥 단부 단변"] and t[1][2] == "(C1)" and t[3][3] == "늑근"
    c = counts([r, b], [1, 2])
    assert c["table"][0]["done"]["기둥"] == 1 and c["table"][1]["done"]["보"] == 1 and c["total"] == 2
    # 접합부 = 보 단부 칸만·늘 단부, 표본 칸·빈 칸 빼기(10/22)
    sch["items"].append({"id": "g1e", "kind": "보", "marks": ["G1"], "floors": "2~5", "position": "단부", "bottom": "4-HD19", "stirrup": "HD10@150", "confirmed": True})
    assert [x["id"] for x in choices(sch, 2, JOINT)] == ["g1e"]
    j = fill_design(norm_row({"kind": JOINT, "floor": 2, "item_id": "g1e", "position": "중앙부"}, 3), sch)
    assert j["position"] == "단부" and j["design"] == ["4EA", "@150"] and member_text(j)[0] == "보 단부(접합부)" and kind_short(j) == "보"
    rows = apply_plan([r, b], default_plan([1, 2]))
    assert len(rows) == 24 and [x["no"] for x in rows][:3] == [1, 2, 3] and rows[0]["kind"] == "보" and rows[0]["floor"] == 1
    assert sum(1 for x in rows if x["floor"] == 1 and x["kind"] == "기둥") == 3 and any(x["id"] == r["id"] for x in rows)
    rows = apply_plan(rows, {"1|기둥": 0, "2|보": 5})
    assert sum(1 for x in rows if x["floor"] == 1 and x["kind"] == "기둥") == 1 and sum(1 for x in rows if x["floor"] == 2 and x["kind"] == "보") == 5
    rep = report_rows(rows)
    assert [x["no"] for x in rep] == [1, 2] and counts(rows, [1, 2])["total"] == 2 and counts(rows, [1, 2])["slots_total"] == len(rows)
    assert len(xlsx([r, b], "철근배근 상태조사")) > 4000


if __name__ == "__main__":
    _selfcheck()
    print("ok")
