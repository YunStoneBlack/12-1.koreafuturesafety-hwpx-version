"""서류 자동화 · 시특법 관급 — 착수계·완수계 엑셀(2026-10-12 2단계, 사용자 "시작"). DB는 모름(값만 받음 — 라우터 server/api/routers/sitok_docs.py).

양식 = 회사 엑셀(바탕화면 `시특법 견본\\서류자동화\\관급_*.xlsx` → data/templates/시특법_착수계_양식.xlsx — 완수계는 산안법 완수계_양식.xlsx, 10/20). 산안법 착수계(server/contract_docs/build.py)와
칸 자리가 같아 그 함수들(갑지 금액 줄·기술자 서류 장·재직증명서)을 같이 쓰고, 다른 점만 여기서:
- 착수계 21장: 갑지·착수계·현장대리인계·[현장대리인 자격증·교육수료증·경력증명서·재직증명서]·책임참여기술자 현황·[참여기술자 4장 × 2명(견본)]·예정공정표·
  보안각서(회사)·[보안각서(사람마다 — 주소·생년월일·성명 + 서명 그림)]. 참여기술자 4장·사람 보안각서는 첫 것을 본으로 사람 수만큼 복사(견본 두 번째 것은 지움).
- 완수계 11장(10/20 형 견본 14쪽 — 붙임 완수내역서 포함): 산안법 완수계 양식 갑지·완수계·완수검사원·청구서 + 결과보고서(연결한 시설물 보고서 표지 그림 —
  없으면 올린 표지) + 회사 서류 6장(4대보험 2·국세·지방세·사업자등록증·통장 — 서류 자동화 설정, 산안법과 같은 것). 예전 이름 "완료계"(10/12~10/20).
- 도장: 양식에 이미 제자리(인쇄 범위 안)에 있음 → 넣기면 그대로(복사한 장에도 같이), 빼기면 지움(산안법은 범위 밖 도장을 새로 놓음 — 여기선 필요 없음).
- 갑지에 문서번호 줄이 없어 산안법 갑지처럼 4행에 "문  서  번  호" + 번호(doc_no — 10/20 형 견본 KFSC26_계약번호_월일, 그 전엔 산안법 KFSC21C_ 규칙).
- 교육수료증 = 정밀안전진단 교육 수료증(시특법 설정 sitok_edu), 자격증·경력증명서 = 서류 자동화 설정(산안법과 같은 그림), 기술등급 = 시특법 등급(건축분야고급기술자 → 고급기술자).
"""
from __future__ import annotations

import datetime
import io
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Side
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image

from server.contract_docs import build, sheet_tools as st

START_TEMPLATE = build.TEMPLATE_DIR / "시특법_착수계_양식.xlsx"
DONE_TEMPLATE = build.DONE_TEMPLATE  # 완수계 = 산안법 완수계 양식 + 손질(10/20 형 견본 — 예전 시특법_완료계_양식.xlsx 4장은 안 씀)
DONE_LIST = ("완수계", "완수검사원", "청구서", "완수내역서", "결과보고서", "4대보험 완납증명서", "국세완납증명서", "지방세완납증명서", "사업자등록증 사본",
             "법인통장 사본")  # 갑지 붙임(견본)
REPORT_AT = 4  # 결과보고서 장 자리(청구서 뒤) — 붙임 완수내역서는 청구서 뒤(contract_docs/attachments.py SLOTS sitok_done)
START_SHEETS = 21


@dataclass
class Person(build.Person):
    sign: str = ""  # 서명 그림(보안각서)


def doc_no(c: build.Contract, day: datetime.date | None) -> str:
    """시특법 관급 문서번호 = KFSC연도2자리_계약번호(끝 차수 "-00" 뺌)_월일 — 10/20 형 완수계 견본 "KFSC26_R26TA02127530_0909"(착수계는 착수일,
    완수계는 완수일). 산안법(build.default_doc_no KFSC21C_계약번호_월일)과 다름."""
    if not c.contract_no:
        return ""
    no = re.sub(r"\s*-\s*\d{1,3}$", "", c.contract_no.strip())
    return f"KFSC{(day or datetime.date.today()):%y}_{no}" + (f"_{day:%m%d}" if day else "")


def greeting(client: str) -> str:
    """"1. 귀 ○의 무궁한 발전을…"의 ○ — 학교면 "교"(견본 양주덕산초), 나머지는 산안법 규칙(build.greeting_word)."""
    return "교" if (client or "").strip().endswith("학교") else build.greeting_word(client)


def grade_text(sitok_grade: str, grade: str) -> str:
    """책임·참여 기술자 현황 기술등급 — 시특법 등급 "건축분야고급기술자" → "고급기술자"(견본), 없으면 산안법 등급."""
    g = re.sub(r"^.*분야\s*", "", (sitok_grade or "").strip())
    return g or grade or ""


def _cover(ws: Worksheet, c: build.Contract, doc_no: str, greet: str, contact: str, contact_cell: str) -> None:
    for cell, v in (("S4", c.client), ("S5", c.title), ("S6", c.contract_no), ("S7", c.amount), ("S8", c.contract_date), ("S9", c.start_date),
                    ("S10", c.end_date)):
        ws[cell] = v
    ws["A4"] = "문서번호"  # 아래 "수 신 자"처럼 A:B 칸에 고르게(양쪽 맞춤)
    ws["A4"].font, ws["A4"].alignment = copy(ws["A5"].font), copy(ws["A5"].alignment)
    if not any("A4" in r for r in ws.merged_cells.ranges):
        ws.merge_cells("A4:B4")
    ws["C4"] = doc_no
    ws["C4"].font, ws["C4"].alignment = copy(ws["C5"].font), copy(ws["C5"].alignment)
    if not any("C4" in r for r in ws.merged_cells.ranges):
        ws.merge_cells("C4:H4")
    # 갑지 로고가 4행까지 내려오면(예전 완료계 양식) 문서번호와 겹침 → 3행을 늘려 4행을 로고 아래로(착수계는 딱 맞아 그대로)
    max_col = st.print_bounds(ws)[0]
    logo_bottom = max((y + h for img in ws._images if img.anchor._from.row < 3 and img.anchor._from.col < max_col
                       for _, y, _, h in [st.anchor_box(ws, img)]), default=0)
    gap = logo_bottom + 3 - st.row_y(ws, 4)
    if gap > 0:
        ws.row_dimensions[3].height = (ws.row_dimensions[3].height or 15) + gap * 72 / 96
    ws["C10"] = f"1. 귀 {greet}의 무궁한 발전을 기원합니다."
    if contact:
        ws[contact_cell] = contact


def _load(path: Path, sheets: int):
    if not path.exists():
        raise FileNotFoundError(f"양식 파일이 없습니다: {path.name} — data/templates에 넣으세요.")
    wb = openpyxl.load_workbook(path)
    if len(wb.worksheets) != sheets:
        raise ValueError(f"{path.name} 시트가 {sheets}장이 아닙니다({len(wb.worksheets)}장) — 양식이 바뀌었으면 server/sitok/gov_docs.py 시트 순서를 맞추세요.")
    return wb


def _copy_sheet(wb, src: Worksheet, pos: int, title: str) -> Worksheet:
    """시트 복사(인쇄 범위 안 그림 — 도장·서명 포함 — 도 같이) → pos 자리."""
    ws = wb.copy_worksheet(src)
    ws.title = title[:31]
    ws.print_area = src.print_area.split("!")[-1]
    ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    ws.sheet_view.view = src.sheet_view.view
    max_col = st.print_bounds(src)[0]
    for img in src._images:
        if img.anchor._from.col < max_col:
            st.add_image_copy(ws, img)
    wb._sheets.remove(ws)
    wb._sheets.insert(pos, ws)
    return ws


def _sign_png(path: str) -> bytes | None:
    """서명 그림 → 흰 바탕을 투명하게 한 PNG("(서명)" 글자 위에 겹쳐도 글자가 보이게)."""
    try:
        with Image.open(path) as im:
            im = im.convert("RGBA")
            im.thumbnail((600, 600))
            px = [(r, g, b, 0 if min(r, g, b) > 200 else a) for r, g, b, a in im.getdata()]
            im.putdata(px)
            buf = io.BytesIO()
            im.save(buf, "PNG")
            return buf.getvalue()
    except Exception:  # noqa: BLE001
        return None


def _pledge(ws: Worksheet, p: Person, sign_box: tuple | None, warnings: list[str], sign_on: bool = True) -> None:
    """사람 보안각서 — B9 주소·생년월일·성명, 서명 그림(본 장의 서명 자리). sign_on=False(도장 빼기 — 10/20 사용자 "대표 도장 뺄 때 이 서명도 같이")면 서명 빈칸."""
    birth = f"{p.birth_date:%Y년 %m월 %d일}" if p.birth_date else ""
    ws["B9"] = f"주        소 : {p.address}\n생 년 월 일 : {birth}\n성        명 : {p.name} (서명)"
    if not sign_on:
        return
    png = _sign_png(p.sign) if p.sign and Path(p.sign).exists() else None
    if png and sign_box:
        st.place_picture(ws, png, sign_box)
    else:
        warnings.append(f"{p.name} 보안각서 — 서명 그림이 없어 빈칸(시특법 설정 탭 기술자 \"서명\"에 올리면 들어감, 아니면 인쇄 뒤 직접 서명).")


def _staff_table(ws: Worksheet, people: list[tuple[str, Person]]) -> None:
    """책임·참여 기술자 현황(5행부터) — 견본은 3줄(7행 아래 굵은 선). 사람 수에 맞춰 줄 모양: 첫 줄 = 5행, 가운데 = 6행, 끝 = 7행(굵은 선)."""
    first, mid, last = ({col: ws[f"{col}{r}"]._style for col in "BCDEF"} for r in (5, 6, 7))
    blank = {col: ws[f"{col}8"]._style for col in "BCDEF"}
    height = ws.row_dimensions[6].height
    n = len(people)
    for i in range(max(n, 3)):
        r = 5 + i
        style = blank if i >= n else last if i == n - 1 else first if i == 0 else mid
        for col in "BCDEF":
            ws[f"{col}{r}"]._style = style[col]
            ws[f"{col}{r}"] = None
        ws.row_dimensions[r].height = height if i < n else ws.row_dimensions[8].height
    for col in "BCDEF":  # 표 아래 줄의 굵은 윗선(견본 3줄 끝) — 줄 수가 바뀌면 빈 줄 위에 선만 남음
        cell = ws[f"{col}{5 + max(n, 3)}"]
        cell.border = Border(left=cell.border.left, right=cell.border.right, bottom=cell.border.bottom)
    for i, (label, p) in enumerate(people):
        r = 5 + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"] = label, "(주)한국미래안전", p.name, p.qualification, p.grade


def _titles(wb) -> None:
    seen: set[str] = set()
    for ws in wb.worksheets:
        name, n = ws.title[:31], 2
        while name in seen:
            name, n = f"{ws.title[:27]} ({n})", n + 1
        seen.add(name)
        if ws.title != name:
            ws.title = name


def build_start(c: build.Contract, doc_no: str, greet: str, contact: str, seal: bool, agent: Person, participants: list[Person], out: Path,
                template: Path = START_TEMPLATE) -> list[str]:
    """착수계 엑셀 → out. 돌려주는 값 = 화면 경고(빠진 서류·서명 등)."""
    warnings: list[str] = []
    wb = _load(template, START_SHEETS)
    ws = wb.worksheets
    cover, start, agent_ws, agent_docs, table = ws[0], ws[1], ws[2], ws[3:7], ws[7]
    part_master, part_extra, plan, pledge_co, pledge_master, pledge_extra = ws[8:12], ws[12:16], ws[16], ws[17], ws[18], ws[19:21]
    _cover(cover, c, doc_no, greet, contact, "B39")
    for s in ws[1:]:
        st.clear_outside_print(s)
    build._amount_line(start, "H7", c.amount)
    agent_ws["H7"], agent_ws["H8"], agent_ws["H9"] = agent.address, agent.name, agent.birth_date
    agent_ws.title = "현장대리인계"
    build._fill_person_sheets(list(agent_docs), agent, "현장대리인", warnings)
    for s, t in zip(agent_docs, ("현장대리인 자격증", "현장대리인 교육수료증", "현장대리인 경력증명서", "현장대리인 재직증명서")):
        s.title = t

    sign = next((img for img in pledge_master._images if not st.is_seal(img) and img.anchor._from.col < st.print_bounds(pledge_master)[0]), None)
    sign_box = st.anchor_box(pledge_master, sign) if sign is not None else None
    st.drop_image(pledge_master, sign)
    for s in list(part_extra) + list(pledge_extra):  # 견본 두 번째 참여기술자·사람 보안각서
        wb.remove(s)

    # 참여기술자: 첫 사람은 본 4장, 둘째부터 본을 복사(마지막 묶음 뒤), 0명이면 본도 지움
    labels = ("자격증", "교육수료증", "경력증명서", "재직증명서")
    groups: list[list[Worksheet]] = []
    if participants:
        groups.append(list(part_master))
        for k, p in enumerate(participants[1:], 2):
            pos = wb.worksheets.index(groups[-1][-1]) + 1
            groups.append([_copy_sheet(wb, s, pos + j, f"참여기술자{k} {labels[j]}") for j, s in enumerate(part_master)])
        for g, p in zip(groups, participants):
            for s, lab in zip(g, labels):
                s.title = f"참여기술자 {lab}({p.name})"[:31]
            build._fill_person_sheets(g, p, "참여기술자", warnings)
    else:
        for s in part_master:
            wb.remove(s)
    _staff_table(table, [("책임기술자", agent)] + [("참여기술자", p) for p in participants])

    # 보안각서: 회사 → 현장대리인 → 참여기술자(본 장을 복사)
    pledge_co.title = "보안각서(회사)"
    pos = wb.worksheets.index(pledge_master)
    copies = [_copy_sheet(wb, pledge_master, pos + 1 + i, f"보안각서({p.name})") for i, p in enumerate(participants)]
    pledge_master.title = f"보안각서({agent.name})"
    for s, p in zip([pledge_master] + copies, [agent] + participants):
        _pledge(s, p, sign_box, warnings, seal)

    if c.start_date is None or c.end_date is None:
        warnings.append("착수일·완수일이 비어 예정공정표 날수가 계산되지 않습니다.")
    plan.title = "예정공정표"
    if not seal:
        for s in wb.worksheets:
            st.remove_seals(s)
    _titles(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    st.fix_dot_numbers(wb)
    wb.save(out)
    return warnings


def build_done(c: build.Contract, doc_no: str, greet: str, contact: str, seal: bool, cover_png: bytes | None, company_docs: dict[str, str], out: Path,
               template: Path = DONE_TEMPLATE, send_date: datetime.date | None = None) -> list[str]:
    """완수계 엑셀 → out(10/20 형 견본 14쪽 "기술지도 완수계처럼 추가 서류"): 산안법 완수계 양식(갑지·완수계·완수검사원·청구서·회사 서류 6장)에
    결과보고서 장(통장 사본 장을 본떠 — 연결한 시설물 보고서 표지)을 청구서 뒤에. 견본과 다른 곳만 손질 — 갑지 붙임 10줄, 완수계 감독관 확인 칸은
    오른쪽 위, 완수검사원 "감독관 경유"·청구서 감독관 칸 없음. 회사 서류 = 서류 자동화 설정(산안법과 같은 것). 발송일·실제준공 = 완수일."""
    warnings: list[str] = []
    wb = _load(template, build.DONE_SHEETS)
    seal_png = st.seal_bytes(wb)
    ws = wb.worksheets
    cover = ws[0]
    end = c.actual_end_date or c.end_date
    settle = c.settle_amount if c.settle_amount is not None else c.amount
    build._fill_cover(cover, c, build.Common(doc_no=doc_no, greeting=greet, contact_name=contact), greet)
    cover["S11"], cover["S12"], cover["H35"] = end, settle, send_date or end or datetime.date.today()  # 발송일 = 창의 발송일 > 실제 준공일 > 준공일(10/22)
    for k in range(11):  # 붙임 C17~C27(양식 11줄 — 기술지도 결과보고서·완료증명서) → 견본 10줄
        cover[f"C{17 + k}"] = f"{k + 1}. {DONE_LIST[k]}" if k < len(DONE_LIST) else None
    for s in (ws[1], ws[2]):  # 완수계·완수검사원 "계약금액"·"정산금액" 줄
        build._amount_line(s, "H7", c.amount)
        build._amount_line(s, "H8", settle)
    build._amount_line(ws[3], "G7", settle)  # 청구서 "청구 금액"
    _inspector_top(ws[1])
    ws[2]["I5"] = None  # 완수검사원 "감독관 경유 (인)"
    _clear_box(ws[3], "B22:F25")  # 청구서 감독관 확인 칸
    for s in ws[1:]:
        st.clear_outside_print(s)
    report = _copy_sheet(wb, ws[9], REPORT_AT, "결과보고서")  # 통장 사본 장 모양(제목 + 큰 그림 칸)
    report["A2"] = "결과보고서"
    st.drop_image(report, st.find_image(report, (209, 64)))  # 원본대조필 칸은 빼고
    old = st.main_image(report)
    if cover_png and old is not None:
        st.replace_image(report, old, cover_png)
    else:
        st.drop_image(report, old)
        warnings.append("결과보고서 — 표지가 없어 제목만 나갑니다(현장 연결한 시설물의 보고서를 만들거나, 표지 PDF·그림을 올리세요).")
    ws = wb.worksheets
    for s, (kind, label) in zip(ws[REPORT_AT + 1:], build.COMPANY_DOCS):
        old = st.main_image(s)
        src = company_docs.get(kind)
        if src and Path(src).exists():
            if old is not None:
                st.replace_image(s, old, src)
        else:
            st.drop_image(s, old)
            warnings.append(f"{label} — 올린 서류가 없어 빈 장으로 나갑니다(서류 자동화 설정에서 올리세요).")
    build._apply_seals(wb, seal, seal_png, text_cells={ws[1].title: ("A25", 41), ws[2].title: ("A25", 45)},
                       extra=lambda: st.place_seal(ws[3], seal_png, st.col_x(ws[3], 11) + 39, st.row_y(ws[3], 20) + st.row_px(ws[3], 20) / 2))
    out.parent.mkdir(parents=True, exist_ok=True)
    st.fix_dot_numbers(wb)
    wb.save(out)
    return warnings


def _clear_box(ws: Worksheet, rng: str) -> None:
    """칸 묶음 비우기 — 글자·선·안쪽 병합."""
    c1, r1, c2, r2 = range_boundaries(rng)
    for m in list(ws.merged_cells.ranges):
        if c1 <= m.min_col and m.max_col <= c2 and r1 <= m.min_row and m.max_row <= r2:
            ws.unmerge_cells(str(m))
    for r in range(r1, r2 + 1):
        for col in range(c1, c2 + 1):
            ws.cell(r, col).value = None
            ws.cell(r, col).border = Border()


def _inspector_top(ws: Worksheet) -> None:
    """완수계 "감독관 확인" 칸 — 양식 왼쪽 아래(B18:F21) → 견본처럼 오른쪽 위(K1:M4, 제목 줄 옆), 제목 "완 수 계"는 그 왼쪽에서 가운데(G2:J3)."""
    head, label, value, sign = (copy(ws[k]._style) for k in ("B18", "B19", "D19", "D21"))
    _clear_box(ws, "B18:F21")
    title, tstyle = ws["A2"].value, copy(ws["A2"]._style)
    ws.unmerge_cells("A2:N3")
    ws["A2"] = None
    ws["G2"], ws["G2"]._style = title, tstyle
    ws.merge_cells("G2:J3")
    for r in range(1, 5):
        ws.row_dimensions[r].height = 18
    for r, text, style in ((1, "감독관 확인", head), (2, "소 속", label), (3, "직 급", label), (4, "성 명", label)):
        ws[f"K{r}"], ws[f"K{r}"]._style = text, style
        for col in "LM":
            ws[f"{col}{r}"]._style = copy(value if r > 1 else head)
    ws["L4"], ws["L4"]._style = "(인)", sign
    thin = Side(style="thin")
    for r in range(1, 5):
        for col in "KLM":
            ws[f"{col}{r}"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells("K1:M1")
    for r in (2, 3, 4):
        ws.merge_cells(f"L{r}:M{r}")


def cover_image(pdf: Path) -> bytes | None:
    """보고서·표지 PDF 첫 쪽 → PNG(결과보고서 장)."""
    import pymupdf

    try:
        with pymupdf.open(pdf) as doc:
            return doc[0].get_pixmap(dpi=110).tobytes("png")
    except Exception:  # noqa: BLE001
        return None
