"""통합 견적 엑셀·PDF(2026-10-29 2단계) — 회사 견적서 양식(data/templates/견적, make_estimate_templates가 숨은 칸 지워 만든 것)에 값만 넣음.

양식(TEMPLATES): 시특법(시특법 정기·정밀·사전사후조사 — 표지 시설물 7칸 × 두 줄 + 줄마다 내역서) · 건축물관리점검(표지 항목 4칸 + 줄마다 내역서·산출근거) ·
안전보건대장(표지 4칸 — 대장 줄마다 대장내역, 이행점검 줄은 이행점검내역) · 기술지도(표지 회당 × 횟수 + 수수료내역, 줄이 하나일 때) · 안전점검(표지 항목 5칸).
값은 전부 글자·숫자로 넣고(엑셀 수식은 지움 — 숨은 칸을 따라가던 것), 금액 한글은 calc.korean_amount(NUMBERSTRING은 LibreOffice에 없음). PDF = LibreOffice(contract_docs.to_pdf)."""
from __future__ import annotations

import datetime
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import range_boundaries

from core.db import BASE_DIR
from server.contract_docs import sheet_tools as st
from server.estimate import calc as C

TEMPLATE_DIR = BASE_DIR / "data" / "templates" / "견적"
KIND_TEMPLATE = {"sitok_regular": "시특법", "sitok_precision": "시특법", "survey": "시특법", "blcm": "건축물관리점검",
                 "health_ledger": "안전보건대장", "impl_check": "안전보건대장", "tech_guide": "기술지도", "site_safety": "안전점검"}
# 항목형 표지 — 칸(항목 줄, 소계 줄), 합계 줄, 견적 대상·조건 첫 줄과 줄 수
ITEM_COVER = {
    "안전점검": {"slots": [(19, 20), (21, 22), (23, 24), (25, 26), (27, 28)], "total": 30, "list": 34, "list_n": 6},
    "건축물관리점검": {"slots": [(19, 20), (21, 22), (23, 24), (25, 26)], "total": 28, "list": 32, "list_n": 10},
    "안전보건대장": {"slots": [(19, 20), (21, 22), (23, 24), (25, 27)], "total": 29, "list": 33, "list_n": 10},
}
SITOK_SLOTS = [24, 26, 28, 30, 32, 34, 36]  # 시설물 칸(두 줄씩)
LEDGER_GRADES = ("기술사", "특급기술자", "고급기술자", "중급기술자", "초급기술자")


# ---------- 칸 도구 ----------
def put(ws, coord: str, value) -> None:
    """칸에 값 — 합친 칸이면 왼쪽 위 칸에."""
    c = ws[coord]
    if isinstance(c, MergedCell):
        for m in ws.merged_cells.ranges:
            if coord in m:
                ws.cell(m.min_row, m.min_col).value = value
                return
        return
    c.value = value


def _bounds(ws) -> tuple[int, int, int, int]:
    area = ws.print_area
    ref = (area[0] if isinstance(area, (list, tuple)) else area).split("!")[-1].replace("$", "")
    return range_boundaries(ref)


def drop_formulas(ws) -> None:
    """채우지 않은 엑셀 수식 지우기(숨은 칸·다른 시트를 따라가던 것 — 남기면 #REF!·NUMBERSTRING이 PDF에 나옴)."""
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c, MergedCell) and isinstance(c.value, str) and c.value.startswith("="):
                c.value = None


def copy_sheet(wb, src, title: str, after=None):
    """시트 복사(그림·인쇄 설정까지 — openpyxl copy_worksheet는 그림·인쇄 범위를 안 옮김)."""
    ws = wb.copy_worksheet(src)
    ws.title = title[:31]
    ws.print_area = src.print_area
    for a in ("page_setup", "print_options", "page_margins"):
        setattr(ws, a, copy(getattr(src, a)))
    ws.sheet_properties.pageSetUpPr = copy(src.sheet_properties.pageSetUpPr)
    ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    ws.sheet_view.view = src.sheet_view.view
    for img in src._images:
        st.add_image_copy(ws, img)
    if after is not None:
        wb._sheets.remove(ws)
        wb._sheets.insert(wb._sheets.index(after) + 1, ws)
    return ws


def _date(s: str | None):
    try:
        return datetime.date.fromisoformat((s or "")[:10])
    except ValueError:
        return datetime.date.today()


def _vat(n: float, rates: dict) -> int:
    return int(round(n * float(rates.get("vat") or 0.1)))


def _won_text(n: int) -> str:
    return f"일금 {C.korean_amount(n)} 원"


# ---------- 표지 ----------
def _head(ws, e: dict) -> None:
    c, p, s = e.get("client") or {}, e.get("project") or {}, e.get("staff") or {}
    put(ws, "F7", e.get("no") or "")
    put(ws, "J7", _date(e.get("issued_on")))
    put(ws, "F8", e.get("title") or "")
    put(ws, "F9", p.get("name") or "")
    put(ws, "F10", c.get("name") or "")
    put(ws, "J10", c.get("ref") or "")
    put(ws, "F11", c.get("phone") or "")
    put(ws, "J11", c.get("email") or "")
    put(ws, "F12", s.get("name") or "")
    put(ws, "H12", s.get("phone") or "")


def cover_sitok(ws, e: dict, outs: list[dict], totals: dict) -> list[str]:
    """시특법 표지 — 시설물 칸 7개 × 두 줄(같은 이름이 이어지면 한 칸의 둘째 줄 — 상반기·하반기), 넘치면 마지막 줄에 "그 외"."""
    notes = []
    c, p = e.get("client") or {}, e.get("project") or {}
    _head(ws, e)
    put(ws, "G16", c.get("name") or "")
    put(ws, "P16", c.get("rep") or "")
    put(ws, "G17", p.get("floors") or "")
    put(ws, "P17", p.get("area") or "")
    put(ws, "G18", p.get("name") or "")
    put(ws, "G19", p.get("address") or "")
    put(ws, "G21", _won_text(totals["total"]))
    put(ws, "P21", None)
    put(ws, "Q21", totals["total"])
    lines = []  # (칸 번호, 줄 0|1, 줄, 결과)
    slot, sub, last = -1, 1, None
    for r, o in zip(e.get("rows") or [], outs):
        name = (r.get("name") or "").strip()
        if slot >= 0 and sub == 0 and name and name == last:
            sub = 1
        else:
            slot, sub = slot + 1, 0
        last = name
        lines.append((slot, sub, r, o))
    if lines and lines[-1][0] >= len(SITOK_SLOTS):
        keep = [x for x in lines if x[0] < len(SITOK_SLOTS) - 1]
        rest = [x for x in lines if x[0] >= len(SITOK_SLOTS) - 1]
        rest_sum = sum(x[3]["supply"] for x in rest)
        keep.append((len(SITOK_SLOTS) - 1, 0, {"name": f"그 외 {len(rest)}건", "check": "", "times": None}, {"unit": None, "supply": rest_sum, "times": None}))
        notes.append(f"시설물이 {len({x[0] for x in lines})}개라 표지 7칸을 넘어 마지막 칸은 '그 외 {len(rest)}건' 합계입니다(내역서는 다 들어감)")
        lines = keep
    for k, row0 in enumerate(SITOK_SLOTS):
        for col in ("D", "E"):
            put(ws, f"{col}{row0}", None)
        for dr in (0, 1):
            for col in ("H", "K", "L", "P"):
                put(ws, f"{col}{row0 + dr}", None)
    for slot, sub, r, o in lines:
        row0 = SITOK_SLOTS[slot]
        put(ws, f"D{row0}", slot + 1)
        put(ws, f"E{row0}", r.get("name") or "")
        put(ws, f"H{row0 + sub}", r.get("check") or "")
        put(ws, f"K{row0 + sub}", o.get("times"))
        put(ws, f"L{row0 + sub}", o.get("unit"))
        put(ws, f"P{row0 + sub}", o.get("supply"))
    put(ws, "L39", totals["supply"])
    put(ws, "P39", "[부가세 별도]")
    return notes


def cover_items(ws, tpl: str, e: dict, outs: list[dict], totals: dict) -> list[str]:
    """항목형 표지(건축물관리점검·안전보건대장·안전점검) — 칸마다 구분·대상·금액 + 소계, 아래 견적 대상·조건."""
    notes = []
    cfg = ITEM_COVER[tpl]
    _head(ws, e)
    put(ws, "G16", _won_text(totals["total"]))
    put(ws, "P16", None)
    put(ws, "Q16", totals["total"])
    rows = list(zip(e.get("rows") or [], outs))
    slots = cfg["slots"]
    if len(rows) > len(slots):
        rest = rows[len(slots) - 1:]
        rows = rows[:len(slots) - 1] + [({"name": f"그 외 {len(rest)}건", "check": ""}, {"supply": sum(o["supply"] for _, o in rest), "times": 1, "unit": None})]
        notes.append(f"항목이 많아 표지 마지막 칸은 '그 외' 합계입니다")
    for top, bot in slots:
        for rr in range(top, bot + 1):
            for col in ("D", "E", "H", "L", "P"):
                put(ws, f"{col}{rr}", None)
    for k, ((r, o), (top, bot)) in enumerate(zip(rows, slots)):
        put(ws, f"D{top}", k + 1)
        put(ws, f"E{top}", r.get("name") or "")
        times = o.get("times") or 1
        what = r.get("check") or ""
        put(ws, f"H{top}", f"{what} {times:g}회".strip() if times and times != 1 else what)
        put(ws, f"L{top}", o.get("supply"))
        put(ws, f"P{top}", r.get("note") or "")
        if bot - top == 2:  # 안전보건대장 이행점검 칸(세 줄): 가운데 = 회당 × 횟수
            put(ws, f"H{top + 1}", f"회당 {int(o.get('unit') or 0):,}원 × {times:g}회" if o.get("unit") else "-")
            put(ws, f"L{top + 1}", None)
        put(ws, f"H{bot}", "소    계")
        put(ws, f"L{bot}", o.get("supply"))
    put(ws, f"L{cfg['total']}", totals["supply"])
    put(ws, f"P{cfg['total']}", "[공급가액] (VAT별도)")
    targets = [f"{i + 1}. {(r.get('name') or '').strip()} {(r.get('check') or '').strip()}".strip() for i, (r, _) in enumerate(zip(e.get("rows") or [], outs))]
    conds = [f"{i + 1}. {x}" for i, x in enumerate((e.get("options") or {}).get("conditions") or [])]
    for i in range(cfg["list_n"]):
        rr = cfg["list"] + i
        put(ws, f"D{rr}", targets[i] if i < len(targets) else None)
        put(ws, f"M{rr}", conds[i] if i < len(conds) else None)
    return notes


def cover_tech(ws, e: dict, outs: list[dict], totals: dict, rates: dict) -> None:
    c, p = e.get("client") or {}, e.get("project") or {}
    o = outs[0]
    _head(ws, e)
    put(ws, "G16", c.get("name") or "")
    put(ws, "P16", c.get("rep") or "")
    put(ws, "G17", None)
    try:
        amt = int(float(str(p.get("amount") or "").replace(",", "")))
    except ValueError:
        amt = p.get("amount") or ""
    put(ws, "H17", amt)
    put(ws, "P17", p.get("period") or "")
    put(ws, "G18", p.get("name") or "")
    put(ws, "G19", p.get("address") or "")
    put(ws, "G21", o.get("times"))
    put(ws, "P21", int(o["unit"]) + _vat(o["unit"], rates))  # 회당 금액(VAT 포함 — 회사 견적서와 같게)
    put(ws, "G22", _won_text(totals["total"]))
    put(ws, "P22", None)
    put(ws, "Q22", totals["total"])
    conds = (e.get("options") or {}).get("conditions") or []
    if conds:
        put(ws, "H29", "\n".join(f"{i + 1}. {x}" for i, x in enumerate(conds)))


# ---------- 계산 시트 ----------
def sheet_sitok(ws, e: dict, r: dict, o: dict, rates: dict) -> None:
    put(ws, "A4", f"문서번호 {e.get('no') or ''}")
    put(ws, "B8", r.get("name") or (e.get("project") or {}).get("name") or "")
    unit = int(o["unit"])
    vat = _vat(unit, rates)
    put(ws, "C10", _won_text(unit + vat))
    put(ws, "F10", unit + vat)
    rate = C.labor_rate(rates, r.get("grade") or "고급기술자", r.get("year")) or 0
    put(ws, "C16", r.get("grade") or "고급기술자")
    put(ws, "D16", r.get("mandays"))
    put(ws, "I16", r.get("field_mandays"))
    put(ws, "I17", 1)
    put(ws, "I18", 1)
    put(ws, "J17", rate)
    put(ws, "K17", o["labor"])
    ex = {x["name"]: x["amount"] for x in o.get("expenses") or []}
    tr, car, hp, rp = r.get("travel") or {}, r.get("car") or {}, r.get("helper") or {}, r.get("report") or {}
    f = lambda v: 0 if v in (None, "") else v  # noqa: E731
    put(ws, "D19", f(tr.get("days"))); put(ws, "G19", f(tr.get("persons"))); put(ws, "I19", f(tr.get("days")) * f(tr.get("persons"))); put(ws, "J19", f(tr.get("price"))); put(ws, "K19", ex.get("여비 및 체제비", 0))
    put(ws, "D20", f(car.get("days"))); put(ws, "G20", f(car.get("count"))); put(ws, "I20", f(car.get("days")) * f(car.get("count"))); put(ws, "J20", f(car.get("price"))); put(ws, "K20", ex.get("차량운영비", 0))
    hp_price = hp.get("price") if hp.get("price") not in (None, "") else rates.get("helper_rate")
    put(ws, "D21", f(hp.get("days"))); put(ws, "G21", f(hp.get("persons"))); put(ws, "I21", f(hp.get("days")) * f(hp.get("persons"))); put(ws, "J21", hp_price); put(ws, "K21", ex.get("현지보조인부", 0))
    put(ws, "I22", f(r.get("danger_pct"))); put(ws, "K22", ex.get("위험수당", 0))
    put(ws, "I23", f(r.get("machine_pct"))); put(ws, "K23", ex.get("기계·기구손료", 0))
    put(ws, "D24", f(rp.get("count"))); put(ws, "J24", f(rp.get("price"))); put(ws, "K24", ex.get("보고서제작", 0))
    put(ws, "K25", sum(v for k, v in ex.items() if k != "추가조사비"))
    put(ws, "K26", ex.get("추가조사비", 0) or None)
    put(ws, "I29", o["overhead_rate"]); put(ws, "K29", o["overhead"])
    put(ws, "I30", o["tech_rate"]); put(ws, "K30", o["tech"])
    put(ws, "K31", o["calc"])
    put(ws, "C32", round(1 - unit / o["calc"], 3) if o["calc"] else 0)
    put(ws, "K32", unit)
    put(ws, "K33", vat)
    put(ws, "K34", unit + vat)
    put(ws, "K35", unit + vat)
    put(ws, "L35", None)
    conds = (e.get("options") or {}).get("conditions") or []
    for i, rr in enumerate((37, 38, 39)):
        put(ws, f"C{rr}", conds[i] if i < len(conds) else None)


def sheet_blcm(ws, ws2, e: dict, r: dict, o: dict, rates: dict) -> None:
    """건축물관리점검 내역서을지·산출근거."""
    year = r.get("year") or rates.get("labor_year")
    rh, rs = C.labor_rate(rates, "기술사", r.get("year")) or 0, C.labor_rate(rates, "초급기술자", r.get("year")) or 0
    heads, staff = (o.get("persons") or [0, 0])
    put(ws, "B2", r.get("name") or (e.get("project") or {}).get("name") or "")
    put(ws, "F6", o["labor"]); put(ws, "M6", f"{year}년 엔지니어링 노임단가 적용")
    put(ws, "B7", "기술사"); put(ws, "C7", heads); put(ws, "E7", rh); put(ws, "F7", int(rh * heads)); put(ws, "M7", "점검책임자(기술사 노임)")
    put(ws, "B8", "초급기술자"); put(ws, "C8", round(staff, 4)); put(ws, "E8", rs); put(ws, "F8", int(rs * staff)); put(ws, "M8", "점검자(초급기술자 노임)")
    put(ws, "C10", o["overhead_rate"]); put(ws, "E10", o["labor"]); put(ws, "F10", o["overhead"])
    put(ws, "D12", o["tech_rate"]); put(ws, "E12", o["labor"] + o["overhead"]); put(ws, "F12", o["tech"])
    put(ws, "C14", 1); put(ws, "E14", o["expense"]); put(ws, "F14", o["expense"])
    put(ws, "F16", o["sum"])
    put(ws, "C18", o.get("adjust")); put(ws, "F18", round(o["sum"] * (o.get("adjust") or 1)))
    put(ws, "F20", o["calc"])
    if int(o["unit"]) != int(o["calc"]):
        put(ws, "A22", "8. 견적금액")
        put(ws, "B22", "네고 반영")
        put(ws, "F22", int(o["unit"]))
    # 산출근거
    age = o.get("age")
    mode = "긴급점검" if r.get("mode") == "urgent" else ("정기점검+구조" if o.get("structure") else "정기점검")
    put(ws2, "J2", _date(e.get("issued_on")))
    put(ws2, "B5", r.get("name") or (e.get("project") or {}).get("name") or "")
    put(ws2, "B7", r.get("area")); put(ws2, "E7", r.get("built_year")); put(ws2, "F7", age)
    put(ws2, "G7", "구조안전 검토필요" if o.get("structure") else ""); put(ws2, "J7", r.get("use") or "")
    put(ws2, "B18", mode); put(ws2, "C18", r.get("area"))
    base_heads = heads - (1 if o.get("structure") and mode != "긴급점검" else 0)
    for col in ("D", "E", "F", "G", "H"):
        put(ws2, f"{col}18", None)
    if mode == "긴급점검":
        put(ws2, "G18", round(heads, 4)); put(ws2, "H18", round(staff, 4))
    else:
        put(ws2, "D18", round(base_heads, 4)); put(ws2, "E18", 1 if o.get("structure") else None); put(ws2, "F18", round(staff, 4))
    age_f = C.blcm_age_factor(age) if age is not None else 1.0
    put(ws2, "B35", r.get("area")); put(ws2, "C35", age); put(ws2, "E35", age_f); put(ws2, "F35", r.get("use") or "")
    put(ws2, "H35", C.BLCM_USE.get((r.get("use") or "").replace(" ", ""), 1.0)); put(ws2, "I35", o.get("adjust"))
    put(ws2, "B40", r.get("area")); put(ws2, "C40", f"{mode} 점검책임자"); put(ws2, "H40", round(heads, 4)); put(ws2, "J40", round(heads, 4)); put(ws2, "J41", round(heads, 4))
    put(ws2, "A45", "건축물"); put(ws2, "B45", r.get("area")); put(ws2, "C45", f"{mode} 점검자"); put(ws2, "H45", round(staff, 4)); put(ws2, "J45", round(staff, 4)); put(ws2, "J46", round(staff, 4))


def sheet_ledger(ws, e: dict, r: dict, o: dict, rates: dict) -> list[str]:
    """공사안전보건대장 내역 — 업무 셋 × 등급 다섯(기술사~초급) 칸. 숙련기술자는 칸이 없어 비고에."""
    notes = []
    put(ws, "A1", f"■ {r.get('name') or '공사안전보건대장'} 작성 용역")
    tasks = r.get("tasks") or []
    for t in range(3):
        top = 6 + 5 * t
        tk = tasks[t] if t < len(tasks) else {"name": "", "lines": []}
        put(ws, f"A{top}", f"1.{t + 1}\n{tk.get('name') or ''}")
        for gi, g in enumerate(LEDGER_GRADES):
            rr = top + gi
            put(ws, f"B{rr}", g)
            put(ws, f"C{rr}", C.labor_rate(rates, g, r.get("year")))
            lns = [ln for ln in tk.get("lines") or [] if ln.get("grade") == g and (ln.get("persons") or 0) and (ln.get("days") or 0)]
            if lns:
                persons = sum(float(ln.get("persons") or 0) for ln in lns)
                days = sum(float(ln.get("persons") or 0) * float(ln.get("days") or 0) for ln in lns) / persons if persons else 0
                put(ws, f"D{rr}", persons); put(ws, f"E{rr}", round(days, 2))
                put(ws, f"F{rr}", (C.labor_rate(rates, g, r.get("year")) or 0) * persons * days)
            else:
                put(ws, f"D{rr}", None); put(ws, f"E{rr}", None); put(ws, f"F{rr}", 0)
        odd = [ln for ln in tk.get("lines") or [] if ln.get("grade") not in LEDGER_GRADES and (ln.get("persons") or 0)]
        if odd:
            notes.append(f"{tk.get('name')}: 숙련기술자 줄은 대장 내역 칸이 없어 금액만 합계에 들어갔습니다")
    if len(tasks) > 3:
        notes.append("업무가 넷 이상이라 대장 내역에는 앞 셋만 칸으로 — 금액은 합계에 다 들어감")
    put(ws, "F4", o["labor"])
    ex = o.get("expenses") or []
    put(ws, "E21", r.get("expense_pct") if r.get("expense_pct") is not None else 0.1); put(ws, "F21", ex[0]["amount"] if ex else 0)
    put(ws, "E23", o["overhead_rate"]); put(ws, "F23", o["overhead"])
    put(ws, "E25", o["tech_rate"]); put(ws, "F25", o["tech"])
    put(ws, "F27", r.get("report") or 0)
    unit = int(o["unit"])
    vat = _vat(unit, rates)
    put(ws, "F28", unit); put(ws, "G28", f"(산출 {int(o['calc']):,}원)" if unit != int(o["calc"]) else "(만원미만 절사)")
    put(ws, "F29", vat)
    put(ws, "F30", unit + vat)
    return notes


def sheet_visit(ws, tpl: str, e: dict, r: dict, o: dict, rates: dict, totals: dict | None) -> None:
    """회당 내역(이행점검내역·수수료내역) — 공사금액 구간 넷 줄, 고른 구간에만 인원·일수."""
    bands = rates.get("visit_bands") or C.VISIT_BANDS_DEFAULT
    chosen = o.get("band")
    for i, b in enumerate(bands[:4]):
        rr = 8 + i
        put(ws, f"A{rr}", b.get("grade"))
        put(ws, f"B{rr}", b.get("label"))
        put(ws, f"C{rr}", b.get("rate"))
        put(ws, f"G{rr}", None)  # 원본 양식에 남은 비고(안전보건전문가 등)
        if b.get("label") == chosen:
            persons, days = r.get("persons") or 1, r.get("days") or 1
            put(ws, f"D{rr}", persons); put(ws, f"E{rr}", days); put(ws, f"F{rr}", o["labor"])
        else:
            put(ws, f"D{rr}", None); put(ws, f"E{rr}", None); put(ws, f"F{rr}", 0)
    put(ws, "F6", o["labor"])
    ex = {x["name"]: x["amount"] for x in o.get("expenses") or []}
    put(ws, "F12", ex.get("직접경비", 0))
    put(ws, "E14", o["overhead_rate"]); put(ws, "F14", o["overhead"])
    put(ws, "E16", o["tech_rate"]); put(ws, "F16", o["tech"])
    put(ws, "F18", ex.get("보고서 제작", 0) or "-")
    calc, unit, times = int(o["calc"]), int(o["unit"]), o.get("times") or 1
    put(ws, "F19", calc)
    if tpl == "수수료내역":
        put(ws, "F20", _vat(calc, rates)); put(ws, "F21", calc + _vat(calc, rates))
        put(ws, "B22", round(unit / calc, 3) if calc else 1)
        per = unit + _vat(unit, rates)
        put(ws, "F22", per); put(ws, "G22", "(VAT 포함)")
        put(ws, "B23", per); put(ws, "D23", times); put(ws, "F23", per * times)
        put(ws, "F24", (totals or {}).get("total", per * times)); put(ws, "G21", f"{times:g}회")
    else:
        put(ws, "A20", "NEGO"); put(ws, "B20", round(1 - unit / calc, 3) if calc else 0); put(ws, "F20", unit)
        put(ws, "F21", _vat(unit, rates)); put(ws, "F22", unit + _vat(unit, rates))
        put(ws, "B23", unit + _vat(unit, rates)); put(ws, "D23", times); put(ws, "F23", (unit + _vat(unit, rates)) * times)
        put(ws, "F24", (unit + _vat(unit, rates)) * times)


# ---------- 한 권 ----------
def build(e: dict, outs: list[dict], totals: dict, rates: dict, dest: Path) -> dict:
    """견적 엑셀 만들기 → {xlsx, notes, sheets}. e = estimates._out 모양(rows·client·project·staff·options …)."""
    kind = e.get("kind")
    tpl = KIND_TEMPLATE.get(kind, "안전점검")
    rows = e.get("rows") or []
    if tpl == "기술지도" and not (len(rows) == 1 and rows[0].get("type") == "visit"):
        tpl = "안전점검"  # 기술지도 표지는 회당 × 횟수 한 줄 — 줄이 여럿이면 항목형으로
    src = TEMPLATE_DIR / f"{tpl}.xlsx"
    if not src.exists():
        raise FileNotFoundError(f"견적 양식이 없습니다({src.name}) — python -m server.scripts.make_estimate_templates <회사 견적서 폴더>")
    wb = openpyxl.load_workbook(src)
    notes: list[str] = []
    cover = wb["견적서"]
    if tpl == "시특법":
        notes += cover_sitok(cover, e, outs, totals)
        base = wb["내역서"]
        sheets = []
        for i, (r, o) in enumerate(zip(rows, outs)):
            if r.get("type") != "sitok":
                continue
            sheets.append((r, o))
        prev = base
        made = []
        for i, (r, o) in enumerate(sheets):
            ws = base if i == 0 else copy_sheet(wb, base, f"내역서 {i + 1}", after=prev)
            made.append((ws, r, o))
            prev = ws
        for i, (ws, r, o) in enumerate(made):
            if len(made) > 1:
                ws.title = f"내역서 {i + 1}"[:31]
            sheet_sitok(ws, e, r, o, rates)
        if not made:
            wb.remove(base)
    elif tpl == "건축물관리점검":
        notes += cover_items(cover, tpl, e, outs, totals)
        b1, b2 = wb["내역서"], wb["산출근거"]
        pairs = [(r, o) for r, o in zip(rows, outs) if r.get("type") == "blcm"]
        prev = b2
        made = []
        for i, (r, o) in enumerate(pairs):
            if i == 0:
                made.append((b1, b2, r, o))
                continue
            w1 = copy_sheet(wb, b1, f"내역서 {i + 1}", after=prev)
            w2 = copy_sheet(wb, b2, f"산출근거 {i + 1}", after=w1)
            made.append((w1, w2, r, o))
            prev = w2
        for w1, w2, r, o in made:
            sheet_blcm(w1, w2, e, r, o, rates)
        if not made:
            wb.remove(b1)
            wb.remove(b2)
    elif tpl == "안전보건대장":
        notes += cover_items(cover, tpl, e, outs, totals)
        led, imp = wb["대장내역"], wb["이행점검내역"]
        eng = [(r, o) for r, o in zip(rows, outs) if r.get("type") == "eng"]
        vis = [(r, o) for r, o in zip(rows, outs) if r.get("type") == "visit"]
        prev = led
        for i, (r, o) in enumerate(eng):
            ws = led if i == 0 else copy_sheet(wb, led, f"대장내역 {i + 1}", after=prev)
            notes += sheet_ledger(ws, e, r, o, rates)
            prev = ws
        prev = imp
        for i, (r, o) in enumerate(vis):
            ws = imp if i == 0 else copy_sheet(wb, imp, f"이행점검내역 {i + 1}", after=prev)
            sheet_visit(ws, "이행점검내역", e, r, o, rates, None)
            prev = ws
        if not eng:
            wb.remove(led)
        if not vis:
            wb.remove(imp)
    elif tpl == "기술지도":
        cover_tech(cover, e, outs, totals, rates)
        sheet_visit(wb["수수료내역"], "수수료내역", e, rows[0], outs[0], rates, totals)
    else:
        notes += cover_items(cover, "안전점검", e, outs, totals)
    for ws in wb.worksheets:
        drop_formulas(ws)
    st.fix_dot_numbers(wb)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return {"xlsx": dest, "notes": notes, "sheets": [ws.title for ws in wb.worksheets], "template": tpl}
