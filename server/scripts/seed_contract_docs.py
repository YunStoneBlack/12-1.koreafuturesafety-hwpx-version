"""착수계·완수계 자료실 첫 채우기(2026-10-06, 한 번만) — 회사 양식(data/templates/*_양식.xlsx)에 이미 들어 있는 것을 DB로 옮긴다.

- 착수계 양식의 현장대리인·참여기술자(재직증명서·현황 표 값 + 자격증·교육수료증·경력증명서 그림) → 기술자 명단
- 완수계 양식의 사업자등록증·통장 사본 그림 → 회사 서류(유효기간 없음)
완납증명서 4장은 옮기지 않는다 — 양식 속 것은 예전에 발급된 것이라 유효기간을 알 수 없음(설정 탭에서 새로 올림).
개인 정보(주소·생년월일)는 코드에 적지 않고 양식에서 읽는다. 같은 이름 기술자가 이미 있으면 건너뜀.

실행: (프로젝트 폴더에서, .env.server의 DATABASE_URL로) python -m server.scripts.seed_contract_docs [--company 1] [--dry-run]
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys

import openpyxl

from core import models_web  # noqa: F401 — 회사 표 정의(외래키)
from core.db import SessionLocal
from core.models_web import SubmitDoc, TechPerson
from server.contract_docs import files
from server.contract_docs import sheet_tools as st
from server.contract_docs.build import COMPANY_DOCS, DONE_TEMPLATE, PERSON_DOCS, START_TEMPLATE

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def _as_date(v) -> datetime.date | None:
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)) and v > 0:  # 엑셀 날짜 일련번호
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
    return None


def _join(text: str) -> datetime.date | None:
    m = re.search(r"(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일", text or "")
    return datetime.date(int(m[1]), int(m[2]), int(m[3])) if m else None


def read_people() -> list[dict]:
    wb = openpyxl.load_workbook(START_TEMPLATE)
    ws = wb.worksheets
    table = ws[7]
    out = []
    for group, row in ((ws[3:7], 5), (ws[8:12], 6)):
        emp = group[3]
        name = (emp["E6"].value or "").strip()
        out.append({
            "name": name,
            "address": (emp["E4"].value or "").strip(),
            "position": (emp["M6"].value or "").strip(),
            "birth_date": _as_date(emp["M7"].value),
            "join_date": _join(emp["B9"].value or ""),
            "qualification": (table[f"E{row}"].value or "").strip(),
            "grade": (table[f"F{row}"].value or "").strip(),
            "images": [st.img_bytes(st.main_image(s)) for s in group[:3]],
        })
    return out


def read_company() -> dict[str, bytes]:
    wb = openpyxl.load_workbook(DONE_TEMPLATE)
    ws = wb.worksheets
    return {kind: st.img_bytes(st.main_image(ws[idx])) for kind, idx in (("biz_reg", 8), ("bankbook", 9))}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--company", type=int, default=1)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    people, company = read_people(), read_company()
    now = datetime.datetime.now()
    with SessionLocal() as db:
        for p in people:
            if db.query(TechPerson).filter(TechPerson.company_id == args.company, TechPerson.name == p["name"]).first():
                print(f"건너뜀(이미 있음): {p['name']}")
                continue
            print(f"기술자 추가: {p['name']} · {p['position']} · {p['grade']} · 생년 {p['birth_date']} · 입사 {p['join_date']}")
            if args.dry_run:
                continue
            row = TechPerson(company_id=args.company, name=p["name"], address=p["address"], position=p["position"],
                             birth_date=p["birth_date"], join_date=p["join_date"], qualification=p["qualification"],
                             grade=p["grade"], active=True, updated_at=now)
            db.add(row)
            db.flush()
            for (kind, label), data in zip(PERSON_DOCS, p["images"]):
                path = files.save_jpeg(data, "x.jpg", files.person_doc_path(row.id, row.name, label))
                db.add(SubmitDoc(company_id=args.company, person_id=row.id, kind=kind, file=str(path), updated_at=now,
                                 updated_by="양식에서 옮김"))
        labels = dict(COMPANY_DOCS)
        for kind, data in company.items():
            exists = db.query(SubmitDoc).filter(SubmitDoc.company_id == args.company, SubmitDoc.person_id.is_(None),
                                                SubmitDoc.kind == kind).first()
            if exists:
                print(f"건너뜀(이미 있음): {labels[kind]}")
                continue
            print(f"회사 서류 추가: {labels[kind]}")
            if args.dry_run:
                continue
            path = files.save_jpeg(data, "x.jpg", files.company_doc_path(labels[kind]))
            db.add(SubmitDoc(company_id=args.company, person_id=None, kind=kind, file=str(path), updated_at=now, updated_by="양식에서 옮김"))
        if not args.dry_run:
            db.commit()


if __name__ == "__main__":
    main()
