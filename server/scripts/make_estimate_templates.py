"""통합 견적 엑셀 양식 만들기(2026-10-29, 한 번 — 회사 견적서가 바뀌면 다시) — 회사 견적서 엑셀(사용자가 준 9개 중 5개)에서 필요한 시트만 남기고
인쇄 범위 밖 칸(오른쪽 "견적서 제출시 변경사항" 입력 칸·노임 도움표·전 고객 정보)과 그림을 비워 data/templates/견적/{양식}.xlsx로. 원본은 원본/에 그대로.

    python -m server.scripts.make_estimate_templates <회사 견적서 폴더>

양식(server/estimate/export.TEMPLATES): 시특법(정기점검 표지 + 내역서) · 건축물관리점검(표지 + 내역서을지 + 산출근거) · 안전보건대장(표지 + 대장 내역 + 이행점검 내역) ·
기술지도(표지 + 수수료내역) · 안전점검(표지). 파일은 시트 이름으로 찾음(받은 파일 이름이 깨져 있어서)."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import range_boundaries

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "data" / "templates" / "견적"
# 양식 이름 → (원본을 알아보는 시트 이름들, 남길 시트 [(원본 시트, 새 이름)])
PLAN = {
    "시특법": (["정기점검", "내역서(1)", "내역서 (4동)"], [("정기점검", "견적서"), ("내역서(1)", "내역서")]),
    "건축물관리점검": (["건축물관리점검", "내역서을지 (1)", "산출근거"], [("건축물관리점검", "견적서"), ("내역서을지 (1)", "내역서"), ("산출근거", "산출근거")]),
    "안전보건대장": (["보건대장 및 이행점검", "안전보건대장", "이행점검"], [("보건대장 및 이행점검", "견적서"), ("안전보건대장", "대장내역"), ("이행점검", "이행점검내역")]),
    "기술지도": (["기술지도", "수수료내역 (3)"], [("기술지도", "견적서"), ("수수료내역 (3)", "수수료내역")]),
    "안전점검": (["안전점검"], [("안전점검", "견적서")]),
}


def bounds(ws) -> tuple[int, int, int, int]:
    area = ws.print_area
    ref = (area[0] if isinstance(area, (list, tuple)) else area).split("!")[-1].replace("$", "")
    return range_boundaries(ref)  # min_col, min_row, max_col, max_row


def clean(ws) -> int:
    """인쇄 범위 밖 값·그림 지우기 → 지운 칸 수."""
    c1, r1, c2, r2 = bounds(ws)
    n = 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c, MergedCell) or c.value is None:
                continue
            if c.column < c1 or c.column > c2 or c.row < r1 or c.row > r2:
                c.value = None
                n += 1
    ws._images = [im for im in ws._images if c1 - 1 <= im.anchor._from.col <= c2 - 1 and r1 - 1 <= im.anchor._from.row <= r2 - 1]
    return n


def find(src_dir: Path, sheets: list[str]) -> Path | None:
    for f in sorted(src_dir.glob("*.xlsx")):
        try:
            names = openpyxl.load_workbook(f, read_only=True).sheetnames
        except Exception:  # noqa: BLE001
            continue
        if all(s in names for s in sheets):
            return f
    return None


def main(src_dir: Path) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "원본").mkdir(exist_ok=True)
    for name, (marks, keep) in PLAN.items():
        f = find(src_dir, marks)
        if f is None:
            print(f"!! {name}: 원본을 못 찾음(시트 {marks})")
            continue
        shutil.copyfile(f, OUT / "원본" / f"{name}.xlsx")
        wb = openpyxl.load_workbook(f)
        want = {s for s, _ in keep}
        for ws in list(wb.worksheets):
            if ws.title not in want:
                wb.remove(ws)
        cleared = 0
        for old, new in keep:
            ws = wb[old]
            cleared += clean(ws)
            ws.title = new
        wb._sheets = [wb[new] for _, new in keep]
        wb.active = 0
        wb.save(OUT / f"{name}.xlsx")
        print(f"{name}: {f.name[:12]} → {[n for _, n in keep]} · 인쇄 범위 밖 {cleared}칸 지움")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(Path(sys.argv[1]))
